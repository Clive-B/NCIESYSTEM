"""Assemble NCIE-002 System Architecture & Technical Design (Production Ready).

Content is authored directly with correct styles (see ncie002_content.py),
so unlike the NCIE-001 pipeline this does not need a reformat-a-baseline
pass — it emits the production-ready docx directly.
"""

from pathlib import Path

from docx import Document

from ncie002_common import (
    DOC_ID,
    DOC_VERSION,
    TOTAL_CHAPTERS,
    add_cover,
    add_governance_block,
    add_toc,
    configure_document,
    render_chapter,
    validate_output,
)
from ncie002_content import CHAPTERS, FRONT_MATTER, set_blocks

BATCH_MODULES = [
    "ncie002_ch01_08",
    "ncie002_ch09_16",
    "ncie002_ch17_24",
    "ncie002_ch25_32",
]


def load_batches():
    import importlib

    loaded = []
    for name in BATCH_MODULES:
        try:
            module = importlib.import_module(name)
        except ImportError:
            continue
        for number, blocks in module.BLOCKS.items():
            set_blocks(number, blocks)
        loaded.append(name)
    return loaded

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "NCIE-002_System_Architecture_Technical_Design_v0_5_IN_DEVELOPMENT.docx"


def add_front_matter_body(document):
    document.add_paragraph("Document Governance", style="NCIE Front Heading")
    add_governance_block(document, FRONT_MATTER["governance_rows"])

    document.add_paragraph("1.4  Naming Convention: ARGUS", style="Heading 2")
    document.add_paragraph(
        "The NCIE Intelligence Assistant is presented to users under the institutional pseudo-name "
        "ARGUS — Adaptive Regulatory Governance & Unified Surveillance. ARGUS is a human-friendly "
        "invocation and presentation identity only; it is not a separate intelligence system, and it "
        "does not replace “NCIE Intelligence Assistant” as the underlying architectural identity used "
        "throughout this specification's component names, interfaces and data model. The full naming "
        "boundary, configurability and audit-identity rules governing ARGUS are defined in Chapter 19 "
        "§19.11. Wherever this document refers to “the Assistant” or “the NCIE Intelligence Assistant,” "
        "ARGUS is that same component under its user-facing name.",
        style="NCIE Body",
    )

    document.add_paragraph("1.5  Chapter Count Correction", style="Heading 2")
    document.add_paragraph(
        "The NCIE Production Documentation Suite register (17 August 2026) recorded NCIE-002 at "
        "20 chapters. During skeletal review, the reviewing architect identified missing "
        "institutional capabilities — most materially active AI collaborative brainstorming "
        "(Human ↔ Human ↔ NCIE) rather than invoked-only Assistant access — which required "
        "new and expanded chapters. The controlled skeleton was revised to Version 0.4 with 32 "
        "chapters. This production-ready edition treats 32 as authoritative; the Suite register "
        "chapter count is stale and should be corrected in a subsequent revision of that register.",
        style="NCIE Body",
    )

    document.add_paragraph("1.6  How Open Review Items Are Handled", style="Heading 2")
    document.add_paragraph(
        "Every chapter of the approved skeleton carried a “Human Review Focus” note flagging "
        "institutional decisions that architecture alone cannot settle — hosting model, "
        "disaster-recovery targets, acceptance authorities, participation-mode defaults, retention "
        "periods, and similar policy calls. Rather than leave these as open questions, this edition "
        "resolves every one of them with a specific, reasoned proposed default, marked inline wherever "
        "it appears with the following convention:",
        style="NCIE Body",
    )
    from ncie002_common import add_proposed
    add_proposed(
        document,
        "This marker means the preceding statement is an architecturally sound default, not a "
        "ratified institutional decision. It must be reviewed and either confirmed or overridden by "
        "the applicable NCIE authority before the affected control is treated as binding.",
    )
    document.add_paragraph(
        "Proposed defaults are chosen to be conservative, auditable and consistent with the mandates "
        "already fixed in NCIE-001 (Human-Primary authority, evidence/provenance-first, AI shall not "
        "approve NCIE for production, etc.). None of them weaken a control already stated as mandatory "
        "in NCIE-001.",
        style="NCIE Body",
    )

    document.add_paragraph("1.7  How to Read This Document", style="Heading 2")
    document.add_paragraph(
        "Each chapter opens with a scope statement, then proceeds through architecture principles, "
        "components and responsibilities, interfaces, data model, event/sequence flows, security "
        "controls, failure modes, non-functional requirements, deployment considerations, acceptance "
        "criteria and traceability back to NCIE-001. Diagrams render as structured tables (flow or "
        "relationship) rather than free-form graphics so the document remains identical across Word, "
        "PDF and any downstream Codex ingestion.",
        style="NCIE Body",
    )

    document.add_paragraph("1.8  Document Amendment History", style="Heading 2")
    for text in [
        "Version: v0.4",
        "Change: Production-ready expansion of the approved 32-chapter skeleton, including the "
        "active-AI-collaborative-brainstorming amendment (Ch.19/20/21/23/24).",
        "Status: APPROVED (historical baseline preserved unchanged).",
        "Version: v0.5",
        "Change: VPF & Dynamic Agent Synthesis Architectural Consistency Amendment (Chapters 1, 19, "
        "20, 24, 28, 31, 32).",
        "Authority: Approved NCIE-001 v1.3 and approved NCIE-007 v1.1.",
        "Scope: Controlled surgical amendment — not a rewrite, restructuring or general architecture "
        "refresh of NCIE-002. The 32-chapter structure is unchanged.",
        "Architecture Impact: Adds Agent Factory, Agent Registry/Capability Registry, Ephemeral "
        "Agent Runtime, VPF, VPF Gate A/B, VPF Profile Registry, and logical Control Plane/Agent "
        "Execution Plane boundaries; extends degraded-mode, acceptance and Codex-handover "
        "architecture accordingly.",
        "Non-Impact: No change to Human-Primary authority, existing domain architecture, Model "
        "Gateway authority, Tool Gateway authority, Evidence governance, Decision authority or "
        "Memory ownership.",
        "Status: v0.5 IN DEVELOPMENT — FOR HUMAN REVIEW OF SURGICAL AMENDMENT. The v0.4 approval "
        "remains historically preserved until v0.5 is reviewed and approved.",
    ]:
        document.add_paragraph(text, style="NCIE Body")

    document.add_page_break()


def main():
    loaded = load_batches()

    document = Document()
    configure_document(document)
    add_cover(document)
    add_front_matter_body(document)
    add_toc(document)

    for chapter in CHAPTERS:
        render_chapter(document, chapter)

    document.core_properties.title = f"{DOC_ID} System Architecture & Technical Design"
    document.core_properties.subject = "NCIE technical architecture — production-ready edition"
    document.core_properties.author = "National Communications Intelligence Ecosystem (NCIE)"
    document.core_properties.category = "Controlled Engineering Document"
    document.core_properties.keywords = (
        "NCIE, system architecture, technical design, communications intelligence, Ghana, Codex"
    )
    document.core_properties.comments = (
        "v0.5: VPF & Dynamic Agent Synthesis Architectural Consistency Amendment (Ch.1,19,20,24,28,"
        "31,32), authority NCIE-001 v1.3 / NCIE-007 v1.1. IN DEVELOPMENT — FOR HUMAN REVIEW."
    )

    document.save(OUTPUT)
    validation = validate_output(OUTPUT)
    print(f"output={OUTPUT}")
    print(f"batches_loaded={loaded}")
    print(f"chapters_written={len(CHAPTERS)}")
    print(f"validation={validation}")


if __name__ == "__main__":
    main()
