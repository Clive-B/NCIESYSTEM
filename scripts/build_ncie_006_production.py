"""Assemble NCIE-006 Memory & Context Architecture Specification (In Development)."""

from pathlib import Path

from docx import Document

from ncie006_common import (
    DOC_ID,
    add_cover,
    add_governance_block,
    add_toc,
    configure_document,
    render_chapter,
    validate_output,
)
from ncie006_content import CHAPTERS, FRONT_MATTER, GOVERNING_INVARIANTS, set_blocks

BATCH_MODULES = [
    "ncie006_ch01_06",
    "ncie006_ch07_12",
    "ncie006_ch13_18",
    "ncie006_ch19_24",
    "ncie006_ch25_30",
    "ncie006_ch31_34",
    "ncie006_ch35_36",
]


def load_batches():
    import importlib

    loaded = []
    for name in BATCH_MODULES:
        try:
            module = importlib.import_module(name)
        except ImportError:
            continue
        for number, blocks in module.BLOCKS.items():
            set_blocks(number, blocks)
        loaded.append(name)
    return loaded


def verify_human_review_register(loaded):
    """Independently re-extract every ('review', ...) block from Chapters 1-34
    at build time and assert the totals match ncie006_ch35_register's own
    computed REGISTER, so a Chapter 35 count mismatch fails the build instead
    of being caught only by inspection (mandatory per handover items 49/110)."""
    import importlib
    from collections import Counter

    from ncie006_ch35_register import BLOCKING_STATUS, REGISTER

    source_modules = [name for name in loaded if name != "ncie006_ch35_36"]
    independent_categories = []
    for name in source_modules:
        module = importlib.import_module(name)
        for _, blocks in module.BLOCKS.items():
            for block in blocks:
                if block[0] != "review":
                    continue
                for entry in block[1]:
                    if isinstance(entry, tuple):
                        independent_categories.append(entry[0])

    register_categories = [it["category"] for it in REGISTER]
    independent_counts = Counter(independent_categories)
    register_counts = Counter(register_categories)

    if len(independent_categories) != len(REGISTER) or independent_counts != register_counts:
        raise AssertionError(
            f"Human Review Register mismatch: independent build-time extraction found "
            f"{len(independent_categories)} items {dict(independent_counts)}, but "
            f"ncie006_ch35_register reports {len(REGISTER)} items {dict(register_counts)}."
        )

    independent_blocking = sum(1 for c in independent_categories if c in BLOCKING_STATUS)
    register_blocking = sum(1 for it in REGISTER if it["category"] in BLOCKING_STATUS)
    if independent_blocking != register_blocking:
        raise AssertionError(
            f"Human Review Register blocking-count mismatch: independent="
            f"{independent_blocking}, register={register_blocking}."
        )

    return {
        "total": len(REGISTER),
        "by_category": dict(register_counts),
        "blocking": register_blocking,
        "non_blocking": len(REGISTER) - register_blocking,
    }


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "NCIE-006_Memory_Context_Architecture_Specification_v1_1_IN_DEVELOPMENT.docx"


def add_front_matter_body(document):
    document.add_paragraph("Document Governance", style="NCIE Front Heading")
    add_governance_block(document, FRONT_MATTER["governance_rows"])

    document.add_paragraph("0.1  Authority Chain", style="Heading 2")
    document.add_paragraph(
        "NCIE-001 requirements → NCIE-002 architecture → NCIE-003 canonical data → NCIE-004 "
        "engineering baseline → NCIE-005 ARGUS behavioral authority → NCIE-006 Memory & Context "
        "architecture. NCIE-006 implements and consolidates Memory/Context requirements established "
        "upstream; it does not reinterpret them, and it does not duplicate NCIE-007 (AI/Agent "
        "Orchestration), NCIE-009 (Security/IAM), NCIE-010 (Evidence/Provenance/Audit), NCIE-011 "
        "(API/Integration), NCIE-014 (Database & Storage Engineering) or NCIE-016 (Testing).",
        style="NCIE Body",
    )

    document.add_paragraph("0.2  Token-Conservation Convention", style="Heading 2")
    document.add_paragraph(
        "Every chapter opens with a short Upstream block citing the specific upstream chapter and "
        "stating only the Memory/Context consequence — never a restatement of the upstream "
        "architecture. Major concepts (Memory, Context, Institutional Knowledge, Decision-Time, "
        "Knowledge Time, Reconciliation, Knowledge Candidate, Context Envelope, Memory Class) are "
        "defined authoritatively once, in the chapter that owns them, and referenced thereafter.",
        style="NCIE Body",
    )

    document.add_paragraph("0.3  Governing Invariants Register", style="Heading 2")
    document.add_paragraph(
        f"These {len(GOVERNING_INVARIANTS)} invariants bind every chapter of this document and are "
        "stated once here rather than re-derived repeatedly:",
        style="NCIE Body",
    )
    for line in GOVERNING_INVARIANTS:
        document.add_paragraph(line, style="List Bullet")

    document.add_paragraph("0.4  Human Review Focus Classification", style="Heading 2")
    document.add_paragraph(
        "Every chapter's Human Review Focus item carries one of six labels: RESOLVED BY UPSTREAM "
        "BASELINE; PROPOSED DESIGN DEFAULT; INSTITUTIONAL CONFIRMATION REQUIRED; SOURCE/WORKFLOW "
        "DISCOVERY REQUIRED; LEGAL/POLICY CONFIRMATION REQUIRED; or SECURITY/SOVEREIGNTY "
        "CONFIRMATION REQUIRED — each additionally marked Blocking or Non-Blocking. Chapter 35 "
        "consolidates every open item, extracted directly from the chapters so the register cannot "
        "drift from its sources (mechanically verified per §35.4).",
        style="NCIE Body",
    )

    document.add_paragraph("0.5  Status Convention", style="Heading 2")
    document.add_paragraph(
        "Skeletal-structure approval authorizes expansion; it does not approve any individual "
        "architectural interpretation or Proposed Design Default. Every chapter in this edition is "
        "stamped IN DEVELOPMENT — FOR HUMAN REVIEW and remains so until explicitly approved. A "
        "technology named from NCIE-004 retains its existing NCIE-004 approval status here — using "
        "it in an NCIE-006 example never promotes it.",
        style="NCIE Body",
    )
    document.add_page_break()


def main():
    loaded = load_batches()
    register_check = verify_human_review_register(loaded)

    document = Document()
    configure_document(document)
    add_cover(document)
    add_front_matter_body(document)
    add_toc(document)

    for chapter in CHAPTERS:
        render_chapter(document, chapter)

    document.core_properties.title = f"{DOC_ID} Memory & Context Architecture Specification"
    document.core_properties.subject = "NCIE Memory & Context architecture — in development"
    document.core_properties.author = "National Communications Intelligence Ecosystem (NCIE)"
    document.core_properties.category = "Controlled Engineering Document"
    document.core_properties.keywords = (
        "NCIE, memory, context, institutional knowledge, ARGUS, communications intelligence, Ghana, Codex"
    )
    document.core_properties.comments = (
        "Expands the approved NCIE-006 skeletal Memory/Context specification (v0.1, 36 chapters). "
        "v1.1: build-time Human Review Register verification; two new invariants on remembered "
        "authorization and prior Decisions. IN DEVELOPMENT — FOR HUMAN REVIEW."
    )

    document.save(OUTPUT)
    validation = validate_output(OUTPUT)
    print(f"output={OUTPUT}")
    print(f"batches_loaded={loaded}")
    print(f"chapters_written={len(CHAPTERS)}")
    print(f"register_check={register_check}")
    print(f"validation={validation}")


if __name__ == "__main__":
    main()
