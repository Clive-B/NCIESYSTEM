from collections import Counter
from datetime import date
from pathlib import Path
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from build_ncie_production_document import (
    CHARCOAL,
    DEEP_GREEN,
    EMERALD,
    FOREST,
    GOLD,
    LIGHT_GREY,
    MID_GREY,
    PALE_GREEN,
    WHITE,
    add_field,
    prevent_row_split,
    set_cell_margins,
    set_cell_shading,
    set_paragraph_bottom_border,
    set_picture_alt_text,
    set_repeat_table_header,
    set_table_borders,
    set_update_fields,
    style_run,
)
from inventory_prd_diagrams import arrow_count, arrow_only, detect as detect_diagram_ranges


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "NCIE-001_Master_PRD_v1_2_EXPANDED_BASELINE.docx"
LOGO = ROOT / "Logo.png"
OUTPUT = ROOT / "NCIE-001_Master_PRD_v1_2_PRODUCTION_READY.docx"

CHAPTER_MARKER_RE = re.compile(
    r"^(?:NCIE-001\s*[\u2014-]\s*)?Chapter\s+(\d+)\s+Expansion$",
    re.I,
)
NUMBERED_L3_RE = re.compile(r"^\d+\.\d+\.\d+\s+")
NUMBERED_L2_RE = re.compile(r"^\d+\.\d+\s+")
NAMED_L3_RE = re.compile(
    r"^(?:Objective|Principle|Phase|Stage|Step|Scenario|Control|Rule|Pattern)\s+\d+\b",
    re.I,
)


def ensure_paragraph_style(document, name, base="Normal"):
    try:
        return document.styles[name]
    except KeyError:
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles[base]
        return style


def set_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_left_border(paragraph, color=GOLD, size="18", space="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    left = p_bdr.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        p_bdr.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), space)
    left.set(qn("w:color"), color)


def remove_paragraph(paragraph):
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def insert_paragraph_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    from docx.text.paragraph import Paragraph

    created = Paragraph(new_p, paragraph._parent)
    if style:
        created.style = style
    if text:
        created.add_run(text)
    return created


def paragraph_all_bold(paragraph):
    runs = [run for run in paragraph.runs if run.text.strip()]
    return bool(runs) and all(run.bold is True for run in runs)


def set_heading_runs(paragraph):
    for run in paragraph.runs:
        run.bold = None
        run.italic = None
        run.font.color.rgb = None
        run.font.name = "Aptos Display"
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Aptos Display")


def set_body_all_regular(paragraph):
    for run in paragraph.runs:
        run.bold = False
        run.font.name = "Aptos"
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Aptos")


def configure_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(10)
    normal.font.bold = False
    normal.font.color.rgb = RGBColor.from_string(CHARCOAL)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    body = ensure_paragraph_style(document, "NCIE Body")
    body.font.name = "Aptos"
    body._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    body.font.size = Pt(10)
    body.font.bold = False
    body.font.color.rgb = RGBColor.from_string(CHARCOAL)
    body.paragraph_format.space_after = Pt(5)
    body.paragraph_format.line_spacing = 1.08
    body.paragraph_format.widow_control = True

    for level, size, color, before, after in (
        (1, 18, DEEP_GREEN, 15, 7),
        (2, 12.5, FOREST, 10, 4),
        (3, 10.8, EMERALD, 8, 3),
    ):
        style = styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    eyebrow = ensure_paragraph_style(document, "NCIE Cover Eyebrow")
    eyebrow.font.name = "Aptos"
    eyebrow._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    eyebrow.font.size = Pt(10)
    eyebrow.font.bold = True
    eyebrow.font.color.rgb = RGBColor.from_string(GOLD)
    eyebrow.paragraph_format.space_after = Pt(4)

    cover_title = ensure_paragraph_style(document, "NCIE Cover Title")
    cover_title.font.name = "Aptos Display"
    cover_title._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
    cover_title.font.size = Pt(27)
    cover_title.font.bold = True
    cover_title.font.color.rgb = RGBColor.from_string(DEEP_GREEN)
    cover_title.paragraph_format.space_after = Pt(3)
    cover_title.paragraph_format.keep_with_next = True

    cover_subtitle = ensure_paragraph_style(document, "NCIE Cover Subtitle")
    cover_subtitle.font.name = "Aptos Display"
    cover_subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
    cover_subtitle.font.size = Pt(20)
    cover_subtitle.font.bold = True
    cover_subtitle.font.color.rgb = RGBColor.from_string(FOREST)
    cover_subtitle.paragraph_format.space_after = Pt(8)

    metadata = ensure_paragraph_style(document, "NCIE Cover Metadata")
    metadata.font.name = "Aptos"
    metadata._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    metadata.font.size = Pt(9.5)
    metadata.font.bold = False
    metadata.font.color.rgb = RGBColor.from_string(CHARCOAL)
    metadata.paragraph_format.space_after = Pt(4)

    chapter = ensure_paragraph_style(document, "NCIE Chapter Label")
    chapter.font.name = "Aptos"
    chapter._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    chapter.font.size = Pt(9)
    chapter.font.bold = True
    chapter.font.color.rgb = RGBColor.from_string(GOLD)
    chapter.paragraph_format.space_after = Pt(3)
    chapter.paragraph_format.keep_with_next = True
    chapter.paragraph_format.page_break_before = True

    status = ensure_paragraph_style(document, "NCIE Status Block")
    status.font.name = "Aptos"
    status._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    status.font.size = Pt(8.5)
    status.font.bold = False
    status.font.color.rgb = RGBColor.from_string(CHARCOAL)
    status.paragraph_format.left_indent = Cm(0.35)
    status.paragraph_format.right_indent = Cm(0.15)
    status.paragraph_format.space_before = Pt(5)
    status.paragraph_format.space_after = Pt(8)
    status.paragraph_format.keep_together = True

    front = ensure_paragraph_style(document, "NCIE Front Heading")
    front.font.name = "Aptos Display"
    front._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
    front.font.size = Pt(18)
    front.font.bold = True
    front.font.color.rgb = RGBColor.from_string(DEEP_GREEN)
    front.paragraph_format.space_after = Pt(8)
    front.paragraph_format.keep_with_next = True

    caption = ensure_paragraph_style(document, "NCIE Figure Caption")
    caption.font.name = "Aptos"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    caption.font.size = Pt(8.5)
    caption.font.bold = False
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(FOREST)
    caption.paragraph_format.space_before = Pt(7)
    caption.paragraph_format.space_after = Pt(3)
    caption.paragraph_format.keep_with_next = True


def configure_page(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.65)
    section.right_margin = Cm(1.65)
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.65)
    section.different_first_page_header_footer = True


def clear_paragraph(paragraph):
    paragraph.clear()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)


def add_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    clear_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("NCIE-001  |  MASTER PRODUCT REQUIREMENTS DOCUMENT  |  VERSION 1.2")
    style_run(run, size=8, color=FOREST, bold=True)
    set_paragraph_bottom_border(p, color=GOLD, size="8", space="3")

    footer = section.footer
    p = footer.paragraphs[0]
    clear_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("CONTROLLED ENGINEERING DOCUMENT     ")
    style_run(run, size=8, color=MID_GREY, bold=False)
    label = p.add_run("Page ")
    style_run(label, size=8, color=MID_GREY, bold=False)
    page = add_field(p, "PAGE")
    style_run(page, size=8, color=MID_GREY, bold=False)
    label = p.add_run(" of ")
    style_run(label, size=8, color=MID_GREY, bold=False)
    pages = add_field(p, "NUMPAGES")
    style_run(pages, size=8, color=MID_GREY, bold=False)


def next_nonempty(paragraphs, start_index):
    for index in range(start_index + 1, len(paragraphs)):
        if paragraphs[index].text.strip():
            return paragraphs[index]
    return None


def previous_is_blank(paragraphs, index):
    return index > 0 and not paragraphs[index - 1].text.strip()


def looks_like_short_heading(paragraphs, index, paragraph):
    text = paragraph.text.strip()
    if not paragraph_all_bold(paragraph):
        return False
    if not previous_is_blank(paragraphs, index):
        return False
    if "\n" in text or len(text) > 110 or len(text.split()) > 14:
        return False
    if text.endswith((".", ";", ":", "?", "!")):
        return False
    if any(token in text for token in ("→", "|", "=", "“", "”")):
        return False
    if text.upper() == text and len(text.split()) > 3:
        return False
    if text.lower().startswith(("chapter status", "nc ie is ", "ncie is ")):
        return False
    return True


def format_cover_and_front_matter(document):
    paragraphs = list(document.paragraphs)
    if len(paragraphs) < 70:
        raise ValueError("Unexpected PRD front-matter structure.")

    logo_p = paragraphs[0].insert_paragraph_before()
    logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_p.paragraph_format.space_after = Pt(8)
    set_paragraph_shading(logo_p, DEEP_GREEN)
    run = logo_p.add_run()
    run.add_picture(str(LOGO), width=Inches(2.65))
    set_picture_alt_text(
        run,
        "NCIE logo",
        "National Communications Intelligence Ecosystem emblem in green, gold, white and red.",
    )

    paragraphs[0].style = "NCIE Cover Eyebrow"
    paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraphs[1].style = "NCIE Cover Title"
    paragraphs[1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_heading_runs(paragraphs[1])
    paragraphs[2].style = "NCIE Cover Subtitle"
    paragraphs[2].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_heading_runs(paragraphs[2])
    set_paragraph_bottom_border(paragraphs[2], color=GOLD, size="18", space="5")

    version_text = paragraphs[3].text.replace("Version 1.1", "Version 1.2")
    paragraphs[3].clear()
    paragraphs[3].style = "NCIE Cover Metadata"
    paragraphs[3].alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraphs[3].add_run(version_text)
    for index in range(4, 10):
        paragraphs[index].style = "NCIE Cover Metadata"
        paragraphs[index].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_body_all_regular(paragraphs[index])

    production_line = insert_paragraph_after(
        paragraphs[9],
        f"Production-ready edition prepared: {date.today().strftime('%d %B %Y')}",
        "NCIE Cover Metadata",
    )
    production_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_body_all_regular(production_line)
    production_line.add_run().add_break(WD_BREAK.PAGE)

    for paragraph in paragraphs[10:24]:
        if not paragraph.text.strip():
            remove_paragraph(paragraph)

    version_note = insert_paragraph_after(
        paragraphs[25],
        "Productionization note: the source filename identifies version 1.2 while the embedded "
        "cover metadata identified Version 1.1. This production-ready edition follows the "
        "filename baseline; custodian confirmation should be recorded before formal release.",
        "NCIE Status Block",
    )
    set_paragraph_shading(version_note, PALE_GREEN)
    set_paragraph_left_border(version_note)
    set_body_all_regular(version_note)

    paragraphs[34].clear()
    paragraphs[34].add_run("Contents")
    paragraphs[34].style = "NCIE Front Heading"
    paragraphs[34].paragraph_format.page_break_before = True
    set_heading_runs(paragraphs[34])

    paragraphs[35].clear()
    paragraphs[35].style = "NCIE Body"
    add_field(paragraphs[35], 'TOC \\o "1-1" \\h \\z \\u')
    for paragraph in paragraphs[36:53]:
        remove_paragraph(paragraph)
    for paragraph in paragraphs[53:69]:
        if not paragraph.text.strip():
            remove_paragraph(paragraph)

    return {
        logo_p,
        *paragraphs[:10],
        production_line,
        version_note,
        paragraphs[34],
        paragraphs[35],
    }


def format_document_paragraphs(document, protected):
    paragraphs = list(document.paragraphs)
    marker_map = {}
    chapter_titles = set()
    for index, paragraph in enumerate(paragraphs):
        match = CHAPTER_MARKER_RE.fullmatch(paragraph.text.strip())
        if not match:
            continue
        number = int(match.group(1))
        if number in marker_map:
            continue
        marker_map[number] = paragraph
        title = next_nonempty(paragraphs, index)
        if title is None:
            raise ValueError(f"Chapter {number} has no title paragraph.")
        chapter_titles.add(title)
        paragraph.style = "NCIE Chapter Label"
        paragraph.paragraph_format.page_break_before = True
        set_heading_runs(paragraph)
        title.style = "Heading 1"
        set_heading_runs(title)
        set_paragraph_bottom_border(title, color=GOLD, size="14", space="5")

    expected = set(range(1, 19))
    if set(marker_map) != expected:
        raise ValueError(f"Chapter marker check failed: found {sorted(marker_map)}")

    counts = Counter()
    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if not text or paragraph in protected or paragraph in marker_map.values() or paragraph in chapter_titles:
            continue
        original_all_bold = paragraph_all_bold(paragraph)

        if paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3"}:
            set_heading_runs(paragraph)
            counts[paragraph.style.name] += 1
            continue
        if text.lower().startswith("chapter status:"):
            paragraph.style = "NCIE Status Block"
            set_paragraph_shading(paragraph, PALE_GREEN)
            set_paragraph_left_border(paragraph)
            set_body_all_regular(paragraph)
            counts["status"] += 1
            continue
        if NUMBERED_L3_RE.match(text):
            paragraph.style = "Heading 3"
            set_heading_runs(paragraph)
            counts["Heading 3"] += 1
            continue
        if NUMBERED_L2_RE.match(text):
            paragraph.style = "Heading 2"
            set_heading_runs(paragraph)
            counts["Heading 2"] += 1
            continue
        if NAMED_L3_RE.match(text) or looks_like_short_heading(paragraphs, index, paragraph):
            paragraph.style = "Heading 3"
            set_heading_runs(paragraph)
            counts["Heading 3"] += 1
            continue

        paragraph.style = "NCIE Body"
        if original_all_bold:
            set_body_all_regular(paragraph)
            counts["debolded_body"] += 1
        else:
            for run in paragraph.runs:
                run.font.name = "Aptos"
                run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Aptos")
        counts["body"] += 1

    return marker_map, counts


def format_tables(document):
    formatted_rows = 0
    for table in document.tables:
        table.autofit = True
        set_table_borders(table, color=LIGHT_GREY, size="4")
        if not table.rows:
            continue
        set_repeat_table_header(table.rows[0])
        for row_index, row in enumerate(table.rows):
            combined = " ".join(cell.text for cell in row.cells)
            if len(combined) < 600:
                prevent_row_split(row)
            fill = DEEP_GREEN if row_index == 0 else (WHITE if row_index % 2 else PALE_GREEN)
            for column_index, cell in enumerate(row.cells):
                set_cell_shading(cell, fill)
                set_cell_margins(cell, top=65, start=75, bottom=65, end=75)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(1)
                    paragraph.paragraph_format.line_spacing = 1.0
                    is_identifier = column_index == 0 and bool(
                        re.match(r"^(?:NCIE|REQ|FR|NFR|ID)[-_]", paragraph.text.strip(), re.I)
                    )
                    for run in paragraph.runs:
                        run.font.name = "Aptos"
                        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Aptos")
                        run.font.size = Pt(8.2)
                        run.font.color.rgb = RGBColor.from_string(
                            WHITE if row_index == 0 else CHARCOAL
                        )
                        run.bold = True if row_index == 0 or is_identifier else False
            formatted_rows += 1
    return formatted_rows


def collect_diagram_blocks(document):
    """Capture diagram paragraph objects before front matter changes shift indices."""
    paragraphs = list(document.paragraphs)
    texts = [paragraph.text.strip() for paragraph in paragraphs]
    ranges = detect_diagram_ranges(texts)
    safe_ranges = []
    for start, end in ranges:
        segment_start = start
        for index in range(start, end + 1):
            if CHAPTER_MARKER_RE.fullmatch(texts[index]):
                if segment_start <= index - 1:
                    safe_ranges.append((segment_start, index - 1))
                segment_start = index + 1
        if segment_start <= end:
            safe_ranges.append((segment_start, end))
    blocks = []
    figure_number = 0
    for start, end in safe_ranges:
        if not any(
            arrow_count(texts[index]) >= 2
            or arrow_only(texts[index])
            or texts[index].count("|") >= 2
            for index in range(start, end + 1)
        ):
            continue
        figure_number += 1
        block_paragraphs = paragraphs[start:end + 1]
        block_texts = [paragraph.text.strip() for paragraph in block_paragraphs]
        blocks.append(
            {
                "number": figure_number,
                "paragraphs": block_paragraphs,
                "texts": block_texts,
            }
        )
    return blocks


def set_table_caption_property(table, value):
    tbl_pr = table._tbl.tblPr
    caption = tbl_pr.find(qn("w:tblCaption"))
    if caption is None:
        caption = OxmlElement("w:tblCaption")
        tbl_pr.append(caption)
    caption.set(qn("w:val"), value)


def set_cell_text(cell, text, *, size=8.2, color=CHARCOAL, bold=False, align=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.style = "NCIE Body"
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    style_run(run, size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell, top=70, start=90, bottom=70, end=90)


def relationship_rows(texts):
    rows = []
    arrow_pattern = re.compile(r"[\u2192\u2190\u2194]+")
    for text in texts:
        for line in text.splitlines():
            parts = [part.strip(" \t-\u2022") for part in arrow_pattern.split(line) if part.strip()]
            if len(parts) == 3 and arrow_count(line) >= 2:
                rows.append(tuple(parts))
    return rows


def diagram_nodes(texts):
    nodes = []
    splitter = re.compile(r"[\u2192\u2190\u2194\u2193\u2191]+")
    for text in texts:
        for line in text.splitlines():
            line = line.strip()
            if not line or arrow_only(line):
                continue
            if arrow_count(line):
                parts = splitter.split(line)
            elif line.count("|") >= 2:
                parts = line.split("|")
            else:
                parts = [line]
            for part in parts:
                value = re.sub(r"^[\s\-\u2022]+", "", part).strip()
                if value and (not nodes or nodes[-1].casefold() != value.casefold()):
                    nodes.append(value)
    return nodes


def add_flow_diagram(document, nodes, figure_number):
    table = document.add_table(rows=0, cols=2)
    table.autofit = True
    set_table_caption_property(table, f"NCIE Figure {figure_number}: process flow")
    set_table_borders(table, color=LIGHT_GREY, size="4")
    for index, node in enumerate(nodes, start=1):
        row = table.add_row()
        prevent_row_split(row)
        number_cell, node_cell = row.cells
        set_cell_shading(number_cell, GOLD)
        set_cell_shading(node_cell, DEEP_GREEN if index % 2 else FOREST)
        set_cell_text(
            number_cell,
            f"{index:02d}",
            size=8.5,
            color=DEEP_GREEN,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        set_cell_text(node_cell, node, size=8.2, color=WHITE, bold=False)
        if index < len(nodes):
            connector = table.add_row()
            prevent_row_split(connector)
            merged = connector.cells[0].merge(connector.cells[1])
            set_cell_shading(merged, PALE_GREEN)
            set_cell_text(
                merged,
                "\u25bc",
                size=6.5,
                color=GOLD,
                bold=True,
                align=WD_ALIGN_PARAGRAPH.CENTER,
            )
            set_cell_margins(merged, top=15, start=20, bottom=15, end=20)
    return table


def add_relationship_diagram(document, rows, figure_number):
    table = document.add_table(rows=1, cols=3)
    table.autofit = True
    set_table_caption_property(table, f"NCIE Figure {figure_number}: relationship model")
    set_table_borders(table, color=LIGHT_GREY, size="4")
    for cell, label in zip(table.rows[0].cells, ("SOURCE", "RELATIONSHIP", "TARGET")):
        set_cell_shading(cell, DEEP_GREEN)
        set_cell_text(cell, label, size=7.5, color=WHITE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table.rows[0])
    for index, (source, relation, target) in enumerate(rows, start=1):
        row = table.add_row()
        prevent_row_split(row)
        fills = (PALE_GREEN, GOLD, PALE_GREEN) if index % 2 else (WHITE, GOLD, WHITE)
        for cell, value, fill in zip(row.cells, (source, relation, target), fills):
            set_cell_shading(cell, fill)
            set_cell_text(
                cell,
                value,
                size=8.0,
                color=DEEP_GREEN if fill == GOLD else CHARCOAL,
                bold=fill == GOLD,
                align=WD_ALIGN_PARAGRAPH.CENTER,
            )
    return table


def format_diagrams(document, blocks):
    converted = 0
    for block in reversed(blocks):
        paragraphs = [p for p in block["paragraphs"] if p._element.getparent() is not None]
        if not paragraphs:
            continue
        texts = block["texts"]
        relationships = relationship_rows(texts)
        nodes = diagram_nodes(texts)
        if len(relationships) >= 2:
            diagram_type = "relationship model"
        elif len(nodes) >= 2:
            diagram_type = "process flow"
        else:
            continue

        anchor = paragraphs[0]
        anchor.clear()
        anchor.style = "NCIE Figure Caption"
        caption_run = anchor.add_run(
            f"Figure {block['number']}. NCIE {diagram_type}."
        )
        caption_run.bold = False
        caption_run.italic = True
        caption_run.font.color.rgb = RGBColor.from_string(FOREST)
        anchor.paragraph_format.keep_with_next = True

        if diagram_type == "relationship model":
            table = add_relationship_diagram(document, relationships, block["number"])
        else:
            table = add_flow_diagram(document, nodes, block["number"])
        anchor._p.addnext(table._tbl)
        for paragraph in paragraphs[1:]:
            remove_paragraph(paragraph)
        converted += 1
    return converted


def validate_output(path):
    document = Document(path)
    paragraphs = list(document.paragraphs)
    chapter_markers = []
    chapter_titles = []
    body_all_bold = []
    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        match = CHAPTER_MARKER_RE.fullmatch(text)
        if match and paragraph.style.name == "NCIE Chapter Label":
            effective_page_break = paragraph.paragraph_format.page_break_before
            if effective_page_break is None:
                effective_page_break = paragraph.style.paragraph_format.page_break_before
            chapter_markers.append((int(match.group(1)), effective_page_break))
            title = next_nonempty(paragraphs, index)
            chapter_titles.append((int(match.group(1)), title.style.name if title else None))
        if paragraph.style.name == "NCIE Body" and paragraph_all_bold(paragraph):
            body_all_bold.append(text[:120])
    if sorted(number for number, _ in chapter_markers) != list(range(1, 19)):
        raise ValueError("Output chapter marker validation failed.")
    if not all(flag is True for _, flag in chapter_markers):
        raise ValueError("One or more chapters do not have a hard page break before them.")
    if not all(style == "Heading 1" for _, style in chapter_titles):
        raise ValueError("One or more chapter titles are not Heading 1.")
    if body_all_bold:
        raise ValueError(f"Unexpected fully bold body paragraphs remain: {body_all_bold[:5]}")
    remaining_diagram_text = [
        paragraph.text[:120]
        for paragraph in paragraphs
        if arrow_only(paragraph.text.strip())
        or arrow_count(paragraph.text) >= 2
        or paragraph.text.count("|") >= 2
    ]
    if remaining_diagram_text:
        raise ValueError(
            f"Unconverted diagram-like paragraphs remain: {remaining_diagram_text[:5]}"
        )
    diagram_tables = [
        table
        for table in document.tables
        if table._tbl.tblPr.find(qn("w:tblCaption")) is not None
        and table._tbl.tblPr.find(qn("w:tblCaption")).get(qn("w:val"), "").startswith("NCIE Figure")
    ]
    return {
        "paragraphs": len(paragraphs),
        "tables": len(document.tables),
        "chapters": len(chapter_markers),
        "diagram_tables": len(diagram_tables),
        "bold_body_paragraphs": len(body_all_bold),
    }


def main():
    document = Document(SOURCE)
    diagram_blocks = collect_diagram_blocks(document)
    configure_styles(document)
    for section in document.sections:
        configure_page(section)
        add_header_footer(section)
    set_update_fields(document)

    protected = format_cover_and_front_matter(document)
    marker_map, counts = format_document_paragraphs(document, protected)
    table_rows = format_tables(document)
    diagram_count = format_diagrams(document, diagram_blocks)

    document.core_properties.title = "NCIE-001 Master Product Requirements Document"
    document.core_properties.subject = "NCIE expanded approved product baseline — production-ready edition"
    document.core_properties.author = "Clive Ebo Barton-Odro"
    document.core_properties.category = "Controlled Engineering Document"
    document.core_properties.keywords = (
        "NCIE, product requirements, communications intelligence, Ghana, governance, Codex"
    )
    document.core_properties.comments = (
        "Production-ready formatting derived from the v1.2 source filename. The source cover "
        "stated Version 1.1; custodian confirmation is required before formal release."
    )

    document.save(OUTPUT)
    validation = validate_output(OUTPUT)
    print(f"output={OUTPUT}")
    print(f"chapter_page_breaks={len(marker_map)}")
    print(f"formatted_table_rows={table_rows}")
    print(f"production_diagrams={diagram_count}")
    print(f"transform_counts={dict(counts)}")
    print(f"validation={validation}")


if __name__ == "__main__":
    main()
