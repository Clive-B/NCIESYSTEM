import json
from pathlib import Path
import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from build_ncie_prd_production import validate_output


ROOT = Path(__file__).resolve().parents[1]
MASTER = ROOT / "NCIE-001_Master_PRD_v1_2_PRODUCTION_READY.docx"
INDEX = ROOT / "prd-page-index.json"
WORKING = ROOT / "NCIE-001_Master_PRD_v1_2_TOC_WORKING.docx"


def add_bookmark(paragraph, name, bookmark_id):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    p_pr = paragraph._p.find(qn("w:pPr"))
    insert_at = 1 if p_pr is not None else 0
    paragraph._p.insert(insert_at, start)
    paragraph._p.insert(insert_at + 1, end)


def remove_existing_ncie_toc_bookmarks(document):
    bookmark_ids = set()
    for node in list(document._element.iter(qn("w:bookmarkStart"))):
        if (node.get(qn("w:name")) or "").startswith("_NCIE001_TOC_"):
            bookmark_ids.add(node.get(qn("w:id")))
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)
    for node in list(document._element.iter(qn("w:bookmarkEnd"))):
        if node.get(qn("w:id")) in bookmark_ids:
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)


def update_toc_paragraph(paragraph, title, page_number, bookmark_name):
    hyperlink = paragraph._p.find(qn("w:hyperlink"))
    if hyperlink is None:
        raise ValueError(f"TOC entry has no hyperlink: {paragraph.text!r}")
    hyperlink.set(qn("w:anchor"), bookmark_name)
    text_nodes = list(hyperlink.iter(qn("w:t")))
    if len(text_nodes) < 2:
        raise ValueError(f"TOC entry has an unexpected field structure: {paragraph.text!r}")
    text_nodes[0].text = title
    text_nodes[-1].text = str(page_number)
    for node in text_nodes[1:-1]:
        node.text = ""
    instruction = hyperlink.find(".//" + qn("w:instrText"))
    if instruction is None:
        raise ValueError(f"TOC entry has no PAGEREF field: {paragraph.text!r}")
    instruction.text = f" PAGEREF {bookmark_name} \\h "


def main():
    page_index = json.loads(INDEX.read_text(encoding="utf-8"))
    if page_index["missing"]:
        raise ValueError(f"PDF page index is incomplete: {page_index['missing']}")
    entries = page_index["entries"]
    document = Document(MASTER)
    remove_existing_ncie_toc_bookmarks(document)

    toc_paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name.lower() == "toc 1"
    ]
    if len(toc_paragraphs) != len(entries):
        raise ValueError(
            f"Expected {len(entries)} TOC entries but found {len(toc_paragraphs)}."
        )

    heading_map = {}
    for paragraph in document.paragraphs:
        text = " ".join(paragraph.text.split())
        if paragraph.style.name == "Heading 1" and text in entries and text not in heading_map:
            heading_map[text] = paragraph
    missing_headings = set(entries) - set(heading_map)
    if missing_headings:
        raise ValueError(f"Missing Heading 1 targets: {sorted(missing_headings)}")

    bookmark_ids = [
        int(node.get(qn("w:id")))
        for node in document._element.iter(qn("w:bookmarkStart"))
        if (node.get(qn("w:id")) or "").isdigit()
    ]
    next_id = max(bookmark_ids, default=0) + 1
    for index, (title, page_number) in enumerate(entries.items(), start=1):
        bookmark_name = f"_NCIE001_TOC_{index:02d}"
        add_bookmark(heading_map[title], bookmark_name, next_id)
        update_toc_paragraph(
            toc_paragraphs[index - 1], title, page_number, bookmark_name
        )
        next_id += 1

    document.core_properties.comments = (
        "Targeted amendments appended to Chapters 14, 15 and 18 on 19 August 2026. "
        "Chapter 14 renumbered as 14.358–14.442; requirement IDs preserved. "
        f"TOC refreshed from the local {page_index['pdf_pages']}-page index."
    )
    document.save(WORKING)
    validation = validate_output(WORKING)
    shutil.copy2(WORKING, MASTER)
    WORKING.unlink()
    print(f"updated_entries={len(entries)}")
    print(f"pdf_pages={page_index['pdf_pages']}")
    print(f"validation={validation}")


if __name__ == "__main__":
    main()
