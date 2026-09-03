from collections import Counter
from copy import deepcopy
from pathlib import Path
import re
import shutil

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

from build_ncie_prd_production import (
    CHARCOAL,
    DEEP_GREEN,
    FOREST,
    GOLD,
    LIGHT_GREY,
    NUMBERED_L2_RE,
    NUMBERED_L3_RE,
    PALE_GREEN,
    WHITE,
    collect_diagram_blocks,
    format_diagrams,
    paragraph_all_bold,
    set_body_all_regular,
    set_heading_runs,
    set_paragraph_left_border,
    set_paragraph_shading,
    validate_output,
)
from build_ncie_production_document import (
    prevent_row_split,
    set_cell_margins,
    set_cell_shading,
    set_repeat_table_header,
    set_table_borders,
)
from replace_prd_chapters import (
    chapter_elements,
    chapter_semantic_text,
    element_text,
    master_chapter_ranges,
    merge_numbering,
    normalize_space,
    renumber_diagrams,
    set_xml_paragraph_text,
)


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "NCIE PRD AMENDMENTS.docx"
MASTER = ROOT / "NCIE-001_Master_PRD_v1_2_PRODUCTION_READY.docx"
WORKING = ROOT / "NCIE-001_Master_PRD_v1_2_AMENDMENT_WORKING.docx"
TARGETS = (14, 15, 18)

AMENDMENT_LABELS = {
    14: "NCIE-001 — Chapter 14 Amendment",
    15: "NCIE-001 — Chapter 15 Amendment",
    18: "NCIE-001 — Chapter 18 Amendment",
}

EXPECTED_REQUIREMENTS = {
    "SEC": range(221, 273),
    "AIA": range(208, 283),
    "GOV": range(241, 333),
}


def ensure_style(document, name, base="Normal"):
    try:
        return document.styles[name]
    except KeyError:
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles[base]
        return style


def configure_amendment_styles(document):
    label = ensure_style(document, "NCIE Amendment Label")
    label.font.name = "Aptos"
    label._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    label.font.size = Pt(9)
    label.font.bold = True
    label.font.color.rgb = RGBColor.from_string(GOLD)
    label.paragraph_format.space_after = Pt(3)
    label.paragraph_format.keep_with_next = True
    label.paragraph_format.page_break_before = True

    title = ensure_style(document, "NCIE Amendment Title")
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
    title.font.size = Pt(16)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(DEEP_GREEN)
    title.paragraph_format.space_after = Pt(7)
    title.paragraph_format.keep_with_next = True


def source_segments(document):
    children = list(document._body._element)
    starts = {}
    label_re = re.compile(r"^NCIE-001\s*[\u2014-]\s*Chapter\s+(14|15|18)\s+Amendment$", re.I)
    for index, child in enumerate(children):
        if child.tag != qn("w:p"):
            continue
        text = normalize_space(element_text(child))
        match = label_re.fullmatch(text)
        if match:
            starts.setdefault(int(match.group(1)), index)
    if set(starts) != set(TARGETS):
        raise ValueError(f"Amendment starts invalid: {starts}")
    ordered = sorted((index, number) for number, index in starts.items())
    segments = {}
    for position, (start, number) in enumerate(ordered):
        end = ordered[position + 1][0] if position + 1 < len(ordered) else len(children)
        while end > start and children[end - 1].tag == qn("w:sectPr"):
            end -= 1
        segments[number] = children[start:end]
    return segments


def renumber_chapter_14_text(text):
    def replacement(match):
        number = int(match.group(1))
        if 366 <= number <= 450:
            return f"14.{number - 8}"
        return match.group(0)

    return re.sub(r"\b14\.(\d{3})\b", replacement, text)


def remove_rendered_page_breaks(element):
    for node in list(element.iter(qn("w:lastRenderedPageBreak"))):
        parent = node.getparent()
        if parent is not None:
            parent.remove(node)


def clean_segments(raw_segments):
    prepared = {}
    for number, raw_elements in raw_segments.items():
        output = []
        seen_15393 = False
        previous_empty = False
        for source_element in raw_elements:
            element = deepcopy(source_element)
            remove_rendered_page_breaks(element)
            if element.tag == qn("w:tbl"):
                for row in list(element.findall(qn("w:tr"))):
                    row_text = element_text(row)
                    if "NCIE-PRD-GOV-287" in row_text and "# NCIE-001" in row_text:
                        element.remove(row)
                for text_node in element.iter(qn("w:t")):
                    value = text_node.text or ""
                    if number == 14:
                        value = renumber_chapter_14_text(value)
                    value = value.replace("APPD", "APPROVED")
                    text_node.text = value
                output.append(element)
                previous_empty = False
                continue

            text = element_text(element)
            normalized = normalize_space(text)
            if normalized == "The# 18.473 Evidence-Gap Acceptance — Continued":
                continue
            if number == 15 and normalized == "15.393 Decision Detection Boundary":
                if seen_15393:
                    continue
                seen_15393 = True
            if number == 14:
                updated = renumber_chapter_14_text(text)
                if updated != text:
                    set_xml_paragraph_text(element, updated)
                    text = updated
                    normalized = normalize_space(text)
            if "APPD" in text:
                updated = text.replace("APPD", "APPROVED")
                set_xml_paragraph_text(element, updated)
                text = updated
                normalized = normalize_space(text)
            if normalized == "NCIE-001 —Chapter 15 Amendment":
                set_xml_paragraph_text(element, AMENDMENT_LABELS[15])
                text = AMENDMENT_LABELS[15]
                normalized = text
            is_empty = not normalized
            if is_empty and previous_empty:
                continue
            output.append(element)
            previous_empty = is_empty
        prepared[number] = output
    return prepared


def append_segments(document, segments):
    body = document._body._element
    for number in reversed(TARGETS):
        ranges = master_chapter_ranges(document)
        _, end = ranges[number]
        children = list(body)
        reference = children[end] if end < len(children) else None
        insert_at = body.index(reference) if reference is not None else len(body) - 1
        for offset, element in enumerate(segments[number]):
            body.insert(insert_at + offset, element)


def amendment_elements(document, number):
    elements = chapter_elements(document, number)
    start = None
    for index, element in enumerate(elements):
        if element.tag == qn("w:p") and normalize_space(element_text(element)) == AMENDMENT_LABELS[number]:
            start = index
            break
    if start is None:
        raise ValueError(f"Chapter {number} amendment label not found after insertion.")
    return elements[start:]


def format_amendment_paragraphs(document, number):
    elements = amendment_elements(document, number)
    paragraphs = [Paragraph(element, document._body) for element in elements if element.tag == qn("w:p")]
    nonempty = [paragraph for paragraph in paragraphs if paragraph.text.strip()]
    label, title = nonempty[0], nonempty[1]
    label.style = "NCIE Amendment Label"
    label.paragraph_format.page_break_before = True
    set_heading_runs(label)
    title.style = "NCIE Amendment Title"
    set_heading_runs(title)

    counts = Counter()
    for paragraph in paragraphs:
        text = paragraph.text.strip()
        if not text or paragraph is label or paragraph is title:
            continue
        original_all_bold = paragraph_all_bold(paragraph)
        if (
            text.startswith("Version:")
            or text.lower().startswith(f"chapter {number} amendment status:")
            or text.lower().startswith("insertion instruction:")
        ):
            paragraph.style = "NCIE Status Block"
            set_paragraph_shading(paragraph, PALE_GREEN)
            set_paragraph_left_border(paragraph)
            set_body_all_regular(paragraph)
            counts["status"] += 1
        elif NUMBERED_L3_RE.match(text):
            paragraph.style = "Heading 3"
            set_heading_runs(paragraph)
            counts["heading3"] += 1
        elif NUMBERED_L2_RE.match(text):
            paragraph.style = "Heading 2"
            set_heading_runs(paragraph)
            counts["heading2"] += 1
        else:
            paragraph.style = "NCIE Body"
            if original_all_bold:
                set_body_all_regular(paragraph)
                counts["debolded"] += 1
            else:
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Aptos")
            counts["body"] += 1
    return counts


def format_amendment_tables(document, number):
    tables = [
        Table(element, document._body)
        for element in amendment_elements(document, number)
        if element.tag == qn("w:tbl")
    ]
    for table in tables:
        table.autofit = True
        set_table_borders(table, color=LIGHT_GREY, size="4")
        if not table.rows:
            continue
        set_repeat_table_header(table.rows[0])
        for row_index, row in enumerate(table.rows):
            prevent_row_split(row)
            fill = DEEP_GREEN if row_index == 0 else (WHITE if row_index % 2 else PALE_GREEN)
            for column_index, cell in enumerate(row.cells):
                set_cell_shading(cell, fill)
                set_cell_margins(cell, top=65, start=75, bottom=65, end=75)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(1)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for run in paragraph.runs:
                        run.font.name = "Aptos"
                        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Aptos")
                        run.font.size = Pt(8.2)
                        run.font.color.rgb = RGBColor.from_string(WHITE if row_index == 0 else CHARCOAL)
                        run.bold = row_index == 0 or (
                            column_index == 0 and run.text.strip().startswith("NCIE-PRD-")
                        )
    return len(tables)


def split_diagram_blocks_at_numbered_headings(blocks):
    """Keep section headings out of diagram conversions.

    The source occasionally places two arrow diagrams back-to-back with the next
    numbered section heading between them.  The generic detector intentionally
    captures broad ranges, so split those ranges here before replacing the
    arrow text with production tables.
    """
    split_blocks = []
    for block in blocks:
        segments = []
        current = []
        for paragraph, text in zip(block["paragraphs"], block["texts"]):
            if NUMBERED_L2_RE.match(text) or NUMBERED_L3_RE.match(text):
                if current:
                    segments.append(current)
                    current = []
                continue
            current.append((paragraph, text))
        if current:
            segments.append(current)
        for segment in segments:
            texts = [text for _, text in segment]
            if not any(
                any(symbol in text for symbol in ("→", "←", "↔", "↓", "↑", "↕"))
                or text.count("|") >= 2
                for text in texts
            ):
                continue
            split_blocks.append(
                {
                    "number": len(split_blocks) + 1,
                    "paragraphs": [paragraph for paragraph, _ in segment],
                    "texts": texts,
                }
            )
    return split_blocks


def validate_amendments(document, original_signatures):
    for number, signature in original_signatures.items():
        current = chapter_semantic_text(document, number)
        position = current.find(AMENDMENT_LABELS[number])
        if position < 0:
            raise ValueError(f"Chapter {number} amendment not found.")
        prefix = current[:position].rstrip()
        if prefix != signature.rstrip():
            raise ValueError(f"Existing Chapter {number} content changed before the amendment.")

    ranges = {14: (358, 442), 15: (309, 470), 18: (366, 536)}
    for chapter, (start, end) in ranges.items():
        counts = Counter()
        for element in amendment_elements(document, chapter):
            if element.tag != qn("w:p"):
                continue
            match = re.match(rf"^{chapter}\.(\d+)\s+", normalize_space(element_text(element)))
            if match:
                counts[int(match.group(1))] += 1
        missing = [value for value in range(start, end + 1) if value not in counts]
        permitted_duplicate = {14: 438, 15: None, 18: 526}[chapter]
        unexpected_duplicates = [
            value for value, count in counts.items()
            if count > 1 and value != permitted_duplicate
        ]
        if missing or unexpected_duplicates:
            raise ValueError(
                f"Chapter {chapter} numbering invalid; missing={missing}, duplicates={unexpected_duplicates}"
            )

    found_ids = []
    for chapter in TARGETS:
        for element in amendment_elements(document, chapter):
            if element.tag != qn("w:tbl"):
                continue
            for row in element.findall(qn("w:tr")):
                text = element_text(row)
                found_ids.extend(re.findall(r"NCIE-PRD-(?:SEC|AIA|GOV)-\d{3}", text))
    counter = Counter(found_ids)
    expected = {
        f"NCIE-PRD-{prefix}-{number:03d}"
        for prefix, numbers in EXPECTED_REQUIREMENTS.items()
        for number in numbers
    }
    if set(counter) != expected or any(count != 1 for count in counter.values()):
        raise ValueError(
            f"Requirement register invalid; missing={sorted(expected-set(counter))}, "
            f"extra={sorted(set(counter)-expected)}, duplicates={[key for key,value in counter.items() if value>1]}"
        )

    amendment_text = "\n".join(
        element_text(element)
        for chapter in TARGETS
        for element in amendment_elements(document, chapter)
    )
    if "APPD" in amendment_text or "# NCIE-001" in amendment_text or "The# 18.473" in amendment_text:
        raise ValueError("Known amendment artifacts remain.")
    if re.search(r"\b14\.(?:44[3-9]|450)\b", amendment_text):
        raise ValueError("Unshifted Chapter 14 amendment numbering remains.")
    return len(expected)


def main():
    source = Document(SOURCE)
    master = Document(MASTER)
    configure_amendment_styles(master)
    original_signatures = {
        number: chapter_semantic_text(master, number) for number in TARGETS
    }
    raw = source_segments(source)
    prepared = clean_segments(raw)
    numbering_mappings = merge_numbering(source, master, prepared)
    append_segments(master, prepared)

    paragraph_counts = {}
    table_count = 0
    for number in TARGETS:
        paragraph_counts[number] = dict(format_amendment_paragraphs(master, number))
        table_count += format_amendment_tables(master, number)

    allowed = {
        element
        for number in TARGETS
        for element in amendment_elements(master, number)
        if element.tag == qn("w:p")
    }
    diagram_blocks = [
        block
        for block in collect_diagram_blocks(master)
        if block["paragraphs"] and all(paragraph._p in allowed for paragraph in block["paragraphs"])
    ]
    diagram_blocks = split_diagram_blocks_at_numbered_headings(diagram_blocks)
    new_diagrams = format_diagrams(master, diagram_blocks)
    total_diagrams = renumber_diagrams(master)
    requirement_count = validate_amendments(master, original_signatures)
    master.core_properties.comments = (
        "Authorized targeted amendments appended to Chapters 14, 15 and 18 on 19 August 2026. "
        "Chapter 14 amendment renumbered to 14.358–14.442; requirement IDs preserved."
    )
    master.save(WORKING)
    validation = validate_output(WORKING)
    shutil.copy2(WORKING, MASTER)
    WORKING.unlink()
    print(f"master={MASTER}")
    print(f"numbering_mappings={numbering_mappings}")
    print(f"amendment_tables={table_count}")
    print(f"new_diagrams={new_diagrams}")
    print(f"total_diagrams={total_diagrams}")
    print(f"new_requirements={requirement_count}")
    print(f"paragraph_counts={paragraph_counts}")
    print(f"validation={validation}")


if __name__ == "__main__":
    main()
