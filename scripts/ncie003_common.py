"""Shared rendering engine for NCIE-003 Data Model & Data Dictionary.

Reuses the exact brand system validated for NCIE-001 and NCIE-002 (fonts,
colours, header/footer conventions, diagram-table rendering), extended with
a classified Human Review Focus register (5 categories) and a distinct
per-chapter status convention ("IN DEVELOPMENT — FOR HUMAN REVIEW" rather
than "APPROVED"), per the authoritative NCIE-003 handover instructions.
"""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from build_ncie_production_document import (
    CHARCOAL,
    DEEP_GREEN,
    EMERALD,
    FOREST,
    GOLD,
    LIGHT_GOLD,
    LIGHT_GREY,
    MID_GREY,
    PALE_GREEN,
    WHITE,
    add_field,
    prevent_row_split,
    set_cell_margins,
    set_cell_shading,
    set_cell_width,
    set_paragraph_bottom_border,
    set_picture_alt_text,
    set_repeat_table_header,
    set_table_borders,
    set_update_fields,
    style_run,
)
from build_ncie_prd_production import (
    add_flow_diagram,
    add_relationship_diagram,
    configure_page,
    configure_styles,
    ensure_paragraph_style,
    set_cell_text,
    set_paragraph_left_border,
    set_paragraph_shading,
    set_table_caption_property,
)

ROOT = Path(__file__).resolve().parents[1]
LOGO = ROOT / "Logo.png"

DOC_ID = "NCIE-003"
DOC_TITLE = "Data Model & Data Dictionary"
DOC_VERSION = "Version 0.4"
DOC_STATUS = "In development — for human review of surgical amendment"
TOTAL_CHAPTERS = 43

REVIEW_CATEGORIES = {
    "RESOLVED": ("RESOLVED BY APPROVED NCIE BASELINE", FOREST),
    "PROPOSED": ("PROPOSED DESIGN DEFAULT", GOLD),
    "INSTITUTIONAL": ("INSTITUTIONAL CONFIRMATION REQUIRED", DEEP_GREEN),
    "SOURCE_DISCOVERY": ("SOURCE/DATA DISCOVERY REQUIRED", EMERALD),
    "LEGAL": ("LEGAL/POLICY CONFIRMATION REQUIRED", CHARCOAL),
}


# ---------------------------------------------------------------------------
# Style setup
# ---------------------------------------------------------------------------

def configure_extra_styles(document):
    styles = document.styles

    bullet = styles["List Bullet"]
    bullet.font.name = "Aptos"
    bullet._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    bullet.font.size = Pt(10)
    bullet.font.bold = False
    bullet.font.color.rgb = RGBColor.from_string(CHARCOAL)
    bullet.paragraph_format.space_after = Pt(4)
    bullet.paragraph_format.line_spacing = 1.08

    try:
        bullet2 = styles["List Bullet 2"]
        bullet2.font.name = "Aptos"
        bullet2._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
        bullet2.font.size = Pt(10)
        bullet2.font.color.rgb = RGBColor.from_string(CHARCOAL)
        bullet2.paragraph_format.space_after = Pt(3)
    except KeyError:
        pass

    proposed = ensure_paragraph_style(document, "NCIE Proposed Block")
    proposed.font.name = "Aptos"
    proposed._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    proposed.font.size = Pt(8.7)
    proposed.paragraph_format.left_indent = Cm(0.35)
    proposed.paragraph_format.right_indent = Cm(0.15)
    proposed.paragraph_format.space_before = Pt(5)
    proposed.paragraph_format.space_after = Pt(8)
    proposed.paragraph_format.keep_together = True

    review = ensure_paragraph_style(document, "NCIE Review Item")
    review.font.name = "Aptos"
    review._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    review.font.size = Pt(8.7)
    review.paragraph_format.left_indent = Cm(0.35)
    review.paragraph_format.right_indent = Cm(0.15)
    review.paragraph_format.space_before = Pt(2)
    review.paragraph_format.space_after = Pt(4)
    review.paragraph_format.keep_together = True

    trace = ensure_paragraph_style(document, "NCIE Trace Block")
    trace.font.name = "Aptos"
    trace._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    trace.font.size = Pt(8.3)
    trace.font.italic = True
    trace.font.color.rgb = RGBColor.from_string(MID_GREY)
    trace.paragraph_format.space_before = Pt(2)
    trace.paragraph_format.space_after = Pt(8)


def configure_document(document):
    configure_styles(document)
    configure_extra_styles(document)
    for section in document.sections:
        configure_page(section)
        add_header_footer(section)
    set_update_fields(document)


def add_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.clear()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(f"{DOC_ID}  |  DATA MODEL & DATA DICTIONARY  |  {DOC_VERSION.upper()}")
    style_run(run, size=8, color=FOREST, bold=True)
    set_paragraph_bottom_border(p, color=GOLD, size="8", space="3")

    footer = section.footer
    p = footer.paragraphs[0]
    p.clear()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("CONTROLLED ENGINEERING DOCUMENT — IN DEVELOPMENT     ")
    style_run(run, size=8, color=MID_GREY, bold=False)
    label = p.add_run("Page ")
    style_run(label, size=8, color=MID_GREY, bold=False)
    page = add_field(p, "PAGE")
    style_run(page, size=8, color=MID_GREY, bold=False)
    label = p.add_run(" of ")
    style_run(label, size=8, color=MID_GREY, bold=False)
    pages = add_field(p, "NUMPAGES")
    style_run(pages, size=8, color=MID_GREY, bold=False)


# ---------------------------------------------------------------------------
# Cover + front matter
# ---------------------------------------------------------------------------

def add_cover(document):
    logo_p = document.add_paragraph()
    logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_p.paragraph_format.space_after = Pt(8)
    set_paragraph_shading(logo_p, DEEP_GREEN)
    run = logo_p.add_run()
    run.add_picture(str(LOGO), width=Inches(2.65))
    set_picture_alt_text(
        run,
        "NCIE logo",
        "National Communications Intelligence Ecosystem emblem in green, gold, white and red.",
    )

    eyebrow = document.add_paragraph(DOC_ID, style="NCIE Cover Eyebrow")
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title = document.add_paragraph(DOC_TITLE, style="NCIE Cover Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = document.add_paragraph(f"{DOC_STATUS} — {DOC_VERSION}", style="NCIE Cover Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_bottom_border(subtitle, color=GOLD, size="18", space="5")

    meta_lines = [
        "Defines how the information required by NCIE-002's architecture is represented, related, governed and persisted",
        "Authoritative canonical data reference for NCIE-004 through NCIE-018 and Codex implementation",
        "Traces to NCIE-001 Master Product Requirements Document v1.2 (requirements authority) and NCIE-002 System Architecture & Technical Design v0.4 (architecture authority)",
        f"Prepared: {date.today().strftime('%d %B %Y')}",
        "Maintained by: National Communications Intelligence Ecosystem (NCIE)",
    ]
    for line in meta_lines:
        p = document.add_paragraph(line, style="NCIE Cover Metadata")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_page_break()


def add_governance_block(document, rows):
    table = document.add_table(rows=0, cols=2)
    table.autofit = False
    set_table_borders(table, color=LIGHT_GREY, size="4")
    for label, value in rows:
        row = table.add_row()
        set_cell_width(row.cells[0], Cm(4.2))
        set_cell_width(row.cells[1], Cm(12.9))
        set_cell_shading(row.cells[0], PALE_GREEN)
        set_cell_shading(row.cells[1], WHITE)
        for cell in row.cells:
            set_cell_margins(cell, top=70, start=110, bottom=70, end=110)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_text(row.cells[0], label, size=8.3, color=FOREST, bold=True)
        set_cell_text(row.cells[1], value, size=8.3, color=CHARCOAL, bold=False)
    document.add_paragraph()
    return table


def add_toc(document):
    heading = document.add_paragraph("Contents", style="NCIE Front Heading")
    heading.paragraph_format.page_break_before = True
    note = document.add_paragraph(style="NCIE Body")
    note.add_run(
        "This table of contents updates automatically in Microsoft Word. If page numbers do not "
        "appear immediately, select the field and choose Update Field."
    )
    toc = document.add_paragraph()
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u')
    document.add_page_break()


# ---------------------------------------------------------------------------
# Tables, diagrams, captions
# ---------------------------------------------------------------------------

_table_counter = {"n": 0}
_figure_counter = {"n": 0}


def add_caption(document, kind, text):
    p = document.add_paragraph(style="NCIE Figure Caption")
    p.add_run(f"{kind}. {text}")
    p.paragraph_format.keep_with_next = True
    return p


def add_data_table(document, headers, rows, caption=None, font_size=8.0, header_size=7.6):
    _table_counter["n"] += 1
    if caption:
        add_caption(document, f"Table {_table_counter['n']}", caption)
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = True
    set_table_borders(table, color=LIGHT_GREY, size="4")
    set_table_caption_property(table, f"NCIE Table {_table_counter['n']}")
    for cell, label in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, DEEP_GREEN)
        set_cell_text(cell, label, size=header_size, color=WHITE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table.rows[0])
    for r_index, row_values in enumerate(rows, start=1):
        row = table.add_row()
        combined_len = sum(len(str(v)) for v in row_values)
        if combined_len < 500:
            prevent_row_split(row)
        fill = WHITE if r_index % 2 else PALE_GREEN
        for cell, value in zip(row.cells, row_values):
            set_cell_shading(cell, fill)
            flagged = str(value).strip().upper() in (
                "SOURCE DISCOVERY REQUIRED",
                "SOURCE/DATA DISCOVERY REQUIRED",
            )
            set_cell_text(
                cell, str(value), size=font_size,
                color=FOREST if flagged else CHARCOAL,
                bold=flagged,
            )
    document.add_paragraph()
    return table


def add_flow(document, nodes, caption=None):
    _figure_counter["n"] += 1
    add_caption(document, f"Figure {_figure_counter['n']}", caption or "NCIE process flow.")
    return add_flow_diagram(document, nodes, _figure_counter["n"])


def add_relationship(document, rows, caption=None):
    _figure_counter["n"] += 1
    add_caption(document, f"Figure {_figure_counter['n']}", caption or "NCIE relationship model.")
    return add_relationship_diagram(document, rows, _figure_counter["n"])


# ---------------------------------------------------------------------------
# Status / proposed / review / trace blocks
# ---------------------------------------------------------------------------

def add_status(document, text):
    p = document.add_paragraph(style="NCIE Status Block")
    set_paragraph_shading(p, PALE_GREEN)
    set_paragraph_left_border(p, color=GOLD)
    run = p.add_run(text)
    style_run(run, size=8.5, color=CHARCOAL, bold=False)
    return p


def add_proposed(document, text):
    p = document.add_paragraph(style="NCIE Proposed Block")
    set_paragraph_shading(p, LIGHT_GOLD)
    set_paragraph_left_border(p, color=GOLD, size="20")
    label = p.add_run("PROPOSED DESIGN DEFAULT.  ")
    style_run(label, size=8.7, color=DEEP_GREEN, bold=True)
    run = p.add_run(text)
    style_run(run, size=8.7, color=CHARCOAL, bold=False)
    return p


def add_upstream_impact(document, text):
    p = document.add_paragraph(style="NCIE Proposed Block")
    set_paragraph_shading(p, LIGHT_GOLD)
    set_paragraph_left_border(p, color=DEEP_GREEN, size="24")
    label = p.add_run("UPSTREAM ARCHITECTURE IMPACT — REVIEW REQUIRED.  ")
    style_run(label, size=8.7, color=DEEP_GREEN, bold=True)
    run = p.add_run(text)
    style_run(run, size=8.7, color=CHARCOAL, bold=False)
    return p


def add_review_item(document, category, text):
    label_text, color = REVIEW_CATEGORIES[category]
    p = document.add_paragraph(style="NCIE Review Item")
    p.paragraph_format.left_indent = Cm(0.35)
    bullet = p.add_run("•  ")
    style_run(bullet, size=8.7, color=color, bold=False)
    label = p.add_run(f"[{label_text}]  ")
    style_run(label, size=8.2, color=color, bold=True)
    run = p.add_run(text)
    style_run(run, size=8.7, color=CHARCOAL, bold=False)
    return p


def add_review_focus(document, items):
    document.add_paragraph("Human Review Focus", style="Heading 2")
    for category, text in items:
        add_review_item(document, category, text)


def add_trace(document, text):
    p = document.add_paragraph(style="NCIE Trace Block")
    run = p.add_run(f"NCIE-001 Traceability: {text}")
    style_run(run, size=8.3, color=MID_GREY, bold=False)
    run.italic = True
    return p


def add_trace002(document, text):
    document.add_paragraph("NCIE-002 Architecture Dependency & Traceability", style="Heading 2")
    document.add_paragraph(text, style="NCIE Body")


# ---------------------------------------------------------------------------
# Chapter + block rendering
# ---------------------------------------------------------------------------

def render_blocks(document, blocks):
    for block in blocks:
        kind = block[0]
        if kind == "h2":
            document.add_paragraph(block[1], style="Heading 2")
        elif kind == "h3":
            document.add_paragraph(block[1], style="Heading 3")
        elif kind == "p":
            document.add_paragraph(block[1], style="NCIE Body")
        elif kind == "bullets":
            for item in block[1]:
                document.add_paragraph(item, style="List Bullet")
        elif kind == "status":
            add_status(document, block[1])
        elif kind == "proposed":
            add_proposed(document, block[1])
        elif kind == "upstream":
            add_upstream_impact(document, block[1])
        elif kind == "trace":
            add_trace(document, block[1])
        elif kind == "trace002":
            add_trace002(document, block[1])
        elif kind == "review":
            add_review_focus(document, block[1])
        elif kind == "flow":
            caption = block[2] if len(block) > 2 else None
            add_flow(document, block[1], caption)
        elif kind == "rel":
            caption = block[2] if len(block) > 2 else None
            add_relationship(document, block[1], caption)
        elif kind == "table":
            headers, rows = block[1], block[2]
            caption = block[3] if len(block) > 3 else None
            add_data_table(document, headers, rows, caption)
        elif kind == "pagebreak":
            document.add_page_break()
        else:
            raise ValueError(f"Unknown block kind: {kind}")


def render_chapter(document, chapter):
    marker = document.add_paragraph(
        f"{DOC_ID} — Chapter {chapter['number']} Expansion", style="NCIE Chapter Label"
    )
    marker.paragraph_format.page_break_before = True
    title = document.add_paragraph(chapter["title"], style="Heading 1")
    set_paragraph_bottom_border(title, color=GOLD, size="14", space="5")
    scope = document.add_paragraph(chapter["scope"], style="NCIE Body")
    scope.runs[0].italic = True
    scope.runs[0].font.color.rgb = RGBColor.from_string(FOREST)
    add_status(document, "Status: IN DEVELOPMENT — FOR HUMAN REVIEW.")
    render_blocks(document, chapter["blocks"])


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------

def validate_output(path):
    document = Document(path)
    paragraphs = list(document.paragraphs)
    chapter_numbers = []
    for paragraph in paragraphs:
        text = paragraph.text.strip()
        if paragraph.style.name == "NCIE Chapter Label" and text.startswith(f"{DOC_ID}"):
            number = int(text.rsplit("Chapter", 1)[1].split("Expansion")[0].strip())
            chapter_numbers.append(number)
    diagram_tables = [
        table
        for table in document.tables
        if table._tbl.tblPr.find(qn("w:tblCaption")) is not None
        and (table._tbl.tblPr.find(qn("w:tblCaption")).get(qn("w:val")) or "").startswith("NCIE Figure")
    ]
    data_tables = [
        table
        for table in document.tables
        if table._tbl.tblPr.find(qn("w:tblCaption")) is not None
        and (table._tbl.tblPr.find(qn("w:tblCaption")).get(qn("w:val")) or "").startswith("NCIE Table")
    ]
    result = {
        "paragraphs": len(paragraphs),
        "tables": len(document.tables),
        "chapters_found": sorted(chapter_numbers),
        "diagram_tables": len(diagram_tables),
        "data_tables": len(data_tables),
    }
    expected = list(range(1, TOTAL_CHAPTERS + 1))
    if sorted(chapter_numbers) != expected:
        missing = sorted(set(expected) - set(chapter_numbers))
        extra = sorted(set(chapter_numbers) - set(expected))
        result["chapter_check"] = f"MISMATCH missing={missing} extra={extra}"
    else:
        result["chapter_check"] = "OK"
    return result
