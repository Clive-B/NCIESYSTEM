"""Assemble NCIE-004 Technology Stack & Engineering Blueprint (In Development)."""

from pathlib import Path

from docx import Document

from ncie004_common import (
    DOC_ID,
    add_cover,
    add_governance_block,
    add_toc,
    configure_document,
    render_chapter,
    validate_output,
)
from ncie004_content import CHAPTERS, FRONT_MATTER, set_blocks

BATCH_MODULES = [
    "ncie004_ch01_06",
    "ncie004_ch07_12",
    "ncie004_ch13_18",
    "ncie004_ch19_24",
    "ncie004_ch25_30",
    "ncie004_ch31_32",
    "ncie004_ch33_36",
]


def load_batches():
    import importlib

    loaded = []
    for name in BATCH_MODULES:
        try:
            module = importlib.import_module(name)
        except ImportError:
            continue
        for number, blocks in module.BLOCKS.items():
            set_blocks(number, blocks)
        loaded.append(name)
    return loaded


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "NCIE-004_Technology_Stack_Engineering_Blueprint_v1_1_IN_DEVELOPMENT.docx"


def add_front_matter_body(document):
    document.add_paragraph("Document Governance", style="NCIE Front Heading")
    add_governance_block(document, FRONT_MATTER["governance_rows"])

    document.add_paragraph("0.1  Authority Chain", style="Heading 2")
    document.add_paragraph(
        "NCIE-001 establishes what NCIE requires. NCIE-002 establishes how the system is "
        "architecturally structured. NCIE-003 establishes how the information that architecture "
        "requires is represented, related, governed and persisted. NCIE-004 selects and governs the "
        "technology families and engineering patterns that implement all three. NCIE-004 implements "
        "upstream authorities; it does not reinterpret them.",
        style="NCIE Body",
    )

    document.add_paragraph("0.2  Token-Conservation Convention", style="Heading 2")
    document.add_paragraph(
        "This document does not reproduce NCIE-002/003 content. Every chapter opens with a short "
        "Upstream block citing the specific upstream chapter and stating only the engineering "
        "consequence for NCIE-004 — never a restatement of the upstream architecture itself. Where a "
        "reader needs the full upstream rationale, the citation is the pointer to go read it there.",
        style="NCIE Body",
    )

    document.add_paragraph("0.3  Technology-Selection Status Classification", style="Heading 2")
    document.add_paragraph(
        "Every technology decision in this document is classified as one of five statuses. None of "
        "the last four is ever silently converted into a confident selection:",
        style="NCIE Body",
    )
    document.add_paragraph("APPROVED UPSTREAM MANDATE — already dictated by NCIE-001/002/003 or another approved institutional decision.", style="List Bullet")
    document.add_paragraph("PROPOSED DESIGN DEFAULT — a technically recommended implementation choice, subject to NCA confirmation.", style="List Bullet")
    document.add_paragraph("INSTITUTIONAL CONFIRMATION REQUIRED — depends on NCA enterprise standards, procurement, existing platforms or operational policy.", style="List Bullet")
    document.add_paragraph("INFRASTRUCTURE DISCOVERY REQUIRED — depends on actual NCA infrastructure, capacity, network, hosting or installed technology.", style="List Bullet")
    document.add_paragraph("SECURITY/SOVEREIGNTY CONFIRMATION REQUIRED — depends on classification, residency, legal, security or data-sovereignty requirements.", style="List Bullet")
    document.add_paragraph(
        "A named product (e.g. PostgreSQL, Kubernetes, Keycloak) is always a Proposed Design Default "
        "until institutionally approved — this document never asserts \"NCIE uses X\" as settled fact.",
        style="NCIE Body",
    )

    document.add_paragraph("0.4  Governance Artifacts and Where They Live", style="Heading 2")
    document.add_paragraph(
        "Chapter 33 is the Technology Registry (governance of what has been approved: version, "
        "owner, license, support/security/lifecycle state). Chapter 35 is the Technology Decision "
        "Matrix (the choosing record: every Decision ID, capability, candidate/default, rationale, "
        "constraints, owner, status) plus the consolidated Human Review Register. Chapter 36 is the "
        "NCIE-002/003 traceability matrix and acceptance/handover gate. Individual chapters reference "
        "Decision IDs (e.g. TD-11-1) rather than repeating this content.",
        style="NCIE Body",
    )

    document.add_paragraph("0.5  Status Convention", style="Heading 2")
    document.add_paragraph(
        "Skeletal-structure approval authorizes expansion; it does not approve any individual "
        "technology recommendation. Every chapter in this edition is stamped IN DEVELOPMENT — FOR "
        "HUMAN REVIEW and remains so until explicitly approved.",
        style="NCIE Body",
    )

    document.add_paragraph("0.6  Hardening Pass (v1.1)", style="Heading 2")
    document.add_paragraph(
        "NCIE-004 v1.0 received a full human review. Disposition: CONDITIONAL PASS — architecturally "
        "sound, targeted technology-governance hardening required before baseline approval. Five "
        "corrections were requested and are applied in this v1.1 edition:",
        style="NCIE Body",
    )
    document.add_paragraph(
        "(1) Every APPROVED UPSTREAM MANDATE technology-table row was audited against what "
        "NCIE-001/002/003 actually mandate; rows naming a specific implementation mechanism, "
        "algorithm or vendor-flavoured pattern beyond what upstream dictates were reclassified "
        "PROPOSED DESIGN DEFAULT (see Ch.33 §33.1A for the audit note and worked examples).",
        style="List Bullet",
    )
    document.add_paragraph(
        "(2) Chapter 35's Technology Decision Matrix and Human Review Register are extracted "
        "programmatically from the chapter tables at build time, not manually counted — the summary "
        "in Ch.35 §35.2 is mechanically correct by construction and cannot drift from the chapters it "
        "consolidates.",
        style="List Bullet",
    )
    document.add_paragraph(
        "(3) Chapter 35 §35.5 now states explicitly that Non-Blocking Proposed Design Defaults may "
        "support downstream design, prototyping, estimation and specification work, but do not "
        "constitute procurement, production or institutional technology approval.",
        style="List Bullet",
    )
    document.add_paragraph(
        "(4) Chapter 33 §33.3 adds the Registry promotion invariant: a Technology Registry entry "
        "never moves to Approved merely because its Proposed Design Default was used downstream; "
        "approval requires an explicit institutional decision against the Decision ID.",
        style="List Bullet",
    )
    document.add_paragraph(
        "(5) Chapter 33's seeded Technology Registry already stated version policies/support windows "
        "rather than pinned exact versions in v1.0; §33.2A now states this convention explicitly so "
        "it is not merely implicit.",
        style="List Bullet",
    )
    document.add_page_break()


def main():
    loaded = load_batches()

    document = Document()
    configure_document(document)
    add_cover(document)
    add_front_matter_body(document)
    add_toc(document)

    for chapter in CHAPTERS:
        render_chapter(document, chapter)

    document.core_properties.title = f"{DOC_ID} Technology Stack & Engineering Blueprint"
    document.core_properties.subject = "NCIE technology stack and engineering blueprint — in development"
    document.core_properties.author = "National Communications Intelligence Ecosystem (NCIE)"
    document.core_properties.category = "Controlled Engineering Document"
    document.core_properties.keywords = (
        "NCIE, technology stack, engineering blueprint, ARGUS, communications intelligence, Ghana, Codex"
    )
    document.core_properties.comments = (
        "Expands the approved NCIE-004 skeletal technology blueprint (v0.1, 36 chapters). Chapters "
        "remain IN DEVELOPMENT — FOR HUMAN REVIEW until explicitly approved; named products are "
        "proposed defaults only."
    )

    document.save(OUTPUT)
    validation = validate_output(OUTPUT)
    print(f"output={OUTPUT}")
    print(f"batches_loaded={loaded}")
    print(f"chapters_written={len(CHAPTERS)}")
    print(f"validation={validation}")


if __name__ == "__main__":
    main()
