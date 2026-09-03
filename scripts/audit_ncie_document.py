from pathlib import Path
import sys

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else ROOT / "NCIE Documentations.docx"


def run_summary(paragraph):
    parts = []
    for run in paragraph.runs:
        text = run.text.replace("\n", "\\n")
        if not text:
            continue
        flags = []
        if run.bold:
            flags.append("B")
        if run.italic:
            flags.append("I")
        if run.underline:
            flags.append("U")
        label = "".join(flags) or "-"
        parts.append(f"[{label}:{text}]")
    return " ".join(parts)


def main():
    document = Document(SOURCE)
    print(f"source={SOURCE}")
    print(f"paragraphs={len(document.paragraphs)}")
    print(f"tables={len(document.tables)}")
    print(f"sections={len(document.sections)}")
    print("\nPARAGRAPHS")
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        print(
            f"{index:03d} | style={paragraph.style.name!r} | "
            f"alignment={paragraph.alignment!s} | text={text!r}"
        )
        print(f"      runs={run_summary(paragraph)}")

    print("\nTABLES")
    for table_index, table in enumerate(document.tables):
        print(f"table={table_index} rows={len(table.rows)} cols={len(table.columns)}")
        for row_index, row in enumerate(table.rows):
            values = [cell.text.replace("\n", " / ").strip() for cell in row.cells]
            print(f"  row={row_index}: {values}")
            bold_cells = []
            for cell_index, cell in enumerate(row.cells):
                bold_text = [
                    run.text
                    for paragraph in cell.paragraphs
                    for run in paragraph.runs
                    if run.bold and run.text.strip()
                ]
                if bold_text:
                    bold_cells.append((cell_index, bold_text))
            if bold_cells:
                print(f"    intentional_bold={bold_cells}")

    print("\nSTYLES IN USE")
    used = sorted({p.style.name for p in document.paragraphs if p.text.strip()})
    for name in used:
        style = document.styles[name]
        print(
            f"{name!r}: font={style.font.name!r}, size={style.font.size}, "
            f"bold={style.font.bold}, italic={style.font.italic}"
        )


if __name__ == "__main__":
    main()
