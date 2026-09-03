from collections import Counter, defaultdict
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "NCIE-001_Master_PRD_v1_2_EXPANDED_BASELINE.docx"


def bold_state(paragraph):
    runs = [run for run in paragraph.runs if run.text.strip()]
    if not runs:
        return "empty"
    flags = [bool(run.bold) for run in runs]
    if all(flags):
        return "all-bold"
    if any(flags):
        return "mixed"
    return "regular"


def main():
    document = Document(SOURCE)
    nonempty = [(i, p) for i, p in enumerate(document.paragraphs) if p.text.strip()]
    style_counts = Counter(p.style.name for _, p in nonempty)
    bold_counts = Counter(bold_state(p) for _, p in nonempty)
    by_style_bold = defaultdict(Counter)
    for _, paragraph in nonempty:
        by_style_bold[paragraph.style.name][bold_state(paragraph)] += 1

    print(f"source={SOURCE}")
    print(f"paragraphs={len(document.paragraphs)}")
    print(f"nonempty_paragraphs={len(nonempty)}")
    print(f"tables={len(document.tables)}")
    print(f"sections={len(document.sections)}")
    print(f"inline_shapes={len(document.inline_shapes)}")
    print(f"core_title={document.core_properties.title!r}")
    print(f"core_subject={document.core_properties.subject!r}")
    print(f"core_author={document.core_properties.author!r}")
    print(f"style_counts={dict(style_counts)}")
    print(f"bold_counts={dict(bold_counts)}")
    print("bold_by_style=")
    for style, counts in by_style_bold.items():
        print(f"  {style!r}: {dict(counts)}")

    print("\nSECTIONS")
    for index, section in enumerate(document.sections):
        print(
            f"{index}: size=({section.page_width},{section.page_height}) "
            f"margins=({section.top_margin},{section.right_margin},"
            f"{section.bottom_margin},{section.left_margin}) "
            f"different_first={section.different_first_page_header_footer}"
        )

    print("\nTABLE SHAPES")
    for index, table in enumerate(document.tables):
        print(f"{index}: rows={len(table.rows)} cols={len(table.columns)}")

    print("\nFIRST 140 NONEMPTY PARAGRAPHS")
    for ordinal, (index, paragraph) in enumerate(nonempty[:140]):
        text = paragraph.text.replace("\n", " / ")
        print(
            f"{ordinal:03d} doc_index={index:04d} style={paragraph.style.name!r} "
            f"bold={bold_state(paragraph)} text={text[:240]!r}"
        )

    print("\nHEADING-LIKE PARAGRAPHS")
    for index, paragraph in nonempty:
        name = paragraph.style.name.lower()
        text = paragraph.text.strip()
        if "heading" in name or text.upper().startswith(("CHAPTER ", "APPENDIX ")):
            print(
                f"doc_index={index:04d} style={paragraph.style.name!r} "
                f"bold={bold_state(paragraph)} text={text[:260]!r}"
            )

    print("\nALL-BOLD NORMAL/BODY SAMPLES")
    shown = 0
    for index, paragraph in nonempty:
        if bold_state(paragraph) != "all-bold":
            continue
        if paragraph.style.name.lower() not in {"normal", "body text"}:
            continue
        print(f"doc_index={index:04d} style={paragraph.style.name!r} text={paragraph.text[:300]!r}")
        shown += 1
        if shown >= 80:
            break


if __name__ == "__main__":
    main()
