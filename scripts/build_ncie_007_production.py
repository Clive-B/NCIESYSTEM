"""Assemble NCIE-007 AI, Agent & Model Orchestration Specification (In Development)."""

from pathlib import Path

from docx import Document

from ncie007_common import (
    DOC_ID,
    add_cover,
    add_governance_block,
    add_toc,
    configure_document,
    render_chapter,
    validate_output,
)
from ncie007_content import CHAPTERS, FRONT_MATTER, GOVERNING_INVARIANTS, set_blocks

BATCH_MODULES = [
    "ncie007_ch01_06",
    "ncie007_ch07_12",
    "ncie007_ch13_18",
    "ncie007_ch19_24",
    "ncie007_ch25_30",
    "ncie007_ch31_34",
    "ncie007_ch35_36",
]


def load_batches():
    import importlib

    loaded = []
    for name in BATCH_MODULES:
        try:
            module = importlib.import_module(name)
        except ImportError:
            continue
        for number, blocks in module.BLOCKS.items():
            set_blocks(number, blocks)
        loaded.append(name)
    return loaded


def verify_human_review_register(loaded):
    """Independently re-extract every ('review', ...) block from Chapters 1-34
    at build time and assert the totals match ncie007_ch35_register's own
    computed REGISTER, so a Chapter 35 count mismatch fails the build instead
    of being caught only by inspection (mandatory per handover items 1311/1369)."""
    import importlib
    from collections import Counter

    from ncie007_ch35_register import BLOCKING_STATUS, REGISTER

    source_modules = [name for name in loaded if name != "ncie007_ch35_36"]
    independent_categories = []
    for name in source_modules:
        module = importlib.import_module(name)
        for _, blocks in module.BLOCKS.items():
            for block in blocks:
                if block[0] != "review":
                    continue
                for entry in block[1]:
                    if isinstance(entry, tuple):
                        independent_categories.append(entry[0])

    register_categories = [it["category"] for it in REGISTER]
    independent_counts = Counter(independent_categories)
    register_counts = Counter(register_categories)

    if len(independent_categories) != len(REGISTER) or independent_counts != register_counts:
        raise AssertionError(
            f"Human Review Register mismatch: independent build-time extraction found "
            f"{len(independent_categories)} items {dict(independent_counts)}, but "
            f"ncie007_ch35_register reports {len(REGISTER)} items {dict(register_counts)}."
        )

    independent_blocking = sum(1 for c in independent_categories if c in BLOCKING_STATUS)
    register_blocking = sum(1 for it in REGISTER if it["category"] in BLOCKING_STATUS)
    if independent_blocking != register_blocking:
        raise AssertionError(
            f"Human Review Register blocking-count mismatch: independent="
            f"{independent_blocking}, register={register_blocking}."
        )

    return {
        "total": len(REGISTER),
        "by_category": dict(register_counts),
        "blocking": register_blocking,
        "non_blocking": len(REGISTER) - register_blocking,
    }


def verify_upstream_traceability(loaded):
    """Confirm every chapter 1-36 carries at least one ('upstream', ...) block
    (mandatory per handover items 1307/1371 — no Orphan NCIE-007 Requirement)."""
    import importlib

    has_upstream = {}
    for name in loaded:
        module = importlib.import_module(name)
        for chnum, blocks in module.BLOCKS.items():
            has_upstream[chnum] = any(b[0] == "upstream" for b in blocks)
    missing = [c for c in range(1, 37) if not has_upstream.get(c)]
    if missing:
        raise AssertionError(f"Chapters missing an upstream block: {missing}")
    return {"chapters_checked": sorted(has_upstream.keys()), "missing": missing}


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "NCIE-007_AI_Agent_Model_Orchestration_Specification_v1_1_IN_DEVELOPMENT.docx"


def add_front_matter_body(document):
    document.add_paragraph("Document Governance", style="NCIE Front Heading")
    add_governance_block(document, FRONT_MATTER["governance_rows"])

    document.add_paragraph("0.1  Authority Chain", style="Heading 2")
    document.add_paragraph(
        "NCIE-001 requirements → NCIE-002 architecture → NCIE-003 canonical data → NCIE-004 "
        "engineering baseline → NCIE-005 ARGUS behavioral authority → NCIE-006 Memory & Context "
        "architecture → NCIE-007 AI/Agent/Model Orchestration. NCIE-007 implements these authorities; "
        "it does not reinterpret them. Where VPF or Dynamic Agent Synthesis introduces an impact on "
        "NCIE-001-006, that impact is recorded in Chapter 36's Upstream Impact Register for a "
        "surgical amendment pass after NCIE-007 is approved — the upstream documents are not amended "
        "during this production cycle.",
        style="NCIE Body",
    )

    document.add_paragraph("0.2  Token-Conservation Convention", style="Heading 2")
    document.add_paragraph(
        "Every chapter opens with a short Upstream block citing the specific upstream chapter and "
        "stating only the orchestration consequence — never a restatement of the upstream "
        "architecture. Human-Primary governance, ARGUS behavior, Memory architecture, canonical data "
        "semantics and the engineering stack are not re-explained; each is referenced once and "
        "applied.",
        style="NCIE Body",
    )

    document.add_paragraph("0.3  Governing Invariants Register", style="Heading 2")
    document.add_paragraph(
        f"These {len(GOVERNING_INVARIANTS)} invariants bind every chapter of this document and are "
        "stated once here rather than re-derived repeatedly:",
        style="NCIE Body",
    )
    for line in GOVERNING_INVARIANTS:
        document.add_paragraph(line, style="List Bullet")

    document.add_paragraph("0.4  VPF and Dynamic Agent Synthesis — Scope Note", style="Heading 2")
    document.add_paragraph(
        "The VPF Validation Framework and Dynamic Agent Synthesis are controlled additions approved "
        "for inclusion in this edition (NCIE-007 Skeletal v0.2). VPF validates AI-generated output; it "
        "does not approve regulatory Decisions, create Findings, authorize blocking or tracking, "
        "approve Institutional Knowledge, or replace human judgment (VPF VALIDATION ≠ HUMAN / "
        "INSTITUTIONAL APPROVAL). Dynamic Agent Synthesis lets ARGUS create bounded task-specific "
        "agents from approved primitives without Codex hand-coding each one; it never creates new "
        "institutional authority, infrastructure or unapproved capability.",
        style="NCIE Body",
    )

    document.add_paragraph("0.5  Human Review Focus Classification", style="Heading 2")
    document.add_paragraph(
        "Every chapter's Human Review Focus item carries one of seven labels: RESOLVED BY UPSTREAM "
        "BASELINE; PROPOSED DESIGN DEFAULT; INSTITUTIONAL CONFIRMATION REQUIRED; SECURITY/SOVEREIGNTY "
        "CONFIRMATION REQUIRED; LEGAL/POLICY CONFIRMATION REQUIRED; ENGINEERING/INFRASTRUCTURE "
        "DISCOVERY REQUIRED; or SOURCE/WORKFLOW DISCOVERY REQUIRED — each additionally marked "
        "Blocking or Non-Blocking, and each stating precisely what it blocks rather than implying the "
        "whole document is blocked. Chapter 35 consolidates every open item, extracted directly from "
        "the chapters so the register cannot drift from its sources (mechanically verified per §35.4).",
        style="NCIE Body",
    )

    document.add_paragraph("0.6  Status Convention", style="Heading 2")
    document.add_paragraph(
        "Skeletal-structure approval authorizes expansion; it does not approve any individual "
        "architectural interpretation or Proposed Design Default. Every chapter in this edition is "
        "stamped IN DEVELOPMENT — FOR HUMAN REVIEW and remains so until explicitly approved. A "
        "technology named from NCIE-004 retains its existing NCIE-004 approval status here — using "
        "it in an NCIE-007 example never promotes it. Document completeness is not document "
        "approval.",
        style="NCIE Body",
    )
    document.add_page_break()


def main():
    loaded = load_batches()
    register_check = verify_human_review_register(loaded)
    traceability_check = verify_upstream_traceability(loaded)

    document = Document()
    configure_document(document)
    add_cover(document)
    add_front_matter_body(document)
    add_toc(document)

    for chapter in CHAPTERS:
        render_chapter(document, chapter)

    document.core_properties.title = f"{DOC_ID} AI, Agent & Model Orchestration Specification"
    document.core_properties.subject = "NCIE AI/Agent/Model Orchestration, VPF and Dynamic Agent Synthesis — in development"
    document.core_properties.author = "National Communications Intelligence Ecosystem (NCIE)"
    document.core_properties.category = "Controlled Engineering Document"
    document.core_properties.keywords = (
        "NCIE, orchestration, ARGUS, agents, VPF, dynamic agent synthesis, model gateway, Ghana, Codex"
    )
    document.core_properties.comments = (
        "Expands the approved NCIE-007 skeletal v0.2 (VPF & Dynamic Agent Update) and NCIE-007-S01 "
        "v0.1. v1.1: clarified that ALL SUBSTANTIVE AI OUTPUT enters VPF by default (profile rigor, "
        "not exemption, varies by risk). IN DEVELOPMENT — FOR HUMAN REVIEW."
    )

    document.save(OUTPUT)
    validation = validate_output(OUTPUT)
    print(f"output={OUTPUT}")
    print(f"batches_loaded={loaded}")
    print(f"chapters_written={len(CHAPTERS)}")
    print(f"register_check={register_check}")
    print(f"traceability_check={traceability_check}")
    print(f"validation={validation}")


if __name__ == "__main__":
    main()
