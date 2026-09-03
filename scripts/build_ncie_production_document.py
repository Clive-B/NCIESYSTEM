from copy import deepcopy
from datetime import date
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "NCIE Documentations.docx"
LOGO = ROOT / "Logo.png"
OUTPUT = ROOT / "NCIE_Documentation_Suite_Production_Ready.docx"

DEEP_GREEN = "071A14"
FOREST = "234F29"
EMERALD = "2A8B46"
GOLD = "D89D26"
LIGHT_GOLD = "F5E8C7"
PALE_GREEN = "EDF4EF"
CHARCOAL = "202621"
MID_GREY = "66706A"
LIGHT_GREY = "E7ECE8"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GREY, size="5"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "nil")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:cantSplit")
    tr_pr.append(marker)


def set_cell_width(cell, width):
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width.twips)))
    tc_w.set(qn("w:type"), "dxa")


def set_paragraph_bottom_border(paragraph, color=GOLD, size="14", space="6"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_field(paragraph, instruction):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend([begin, text, separate, end])
    return run


def set_picture_alt_text(run, title, description):
    for doc_pr in run._r.xpath(".//wp:docPr"):
        doc_pr.set("title", title)
        doc_pr.set("descr", description)


def set_update_fields(document):
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def style_run(run, *, size=None, color=None, bold=None, name="Aptos"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def configure_styles(document):
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(CHARCOAL)
    normal.font.bold = False
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    body = styles.add_style("NCIE Body", WD_STYLE_TYPE.PARAGRAPH)
    body.base_style = normal
    body.font.name = "Aptos"
    body._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    body.font.size = Pt(10.5)
    body.font.bold = False
    body.paragraph_format.space_after = Pt(7)
    body.paragraph_format.line_spacing = 1.12

    for level, size, color, before, after in (
        (1, 18, DEEP_GREEN, 16, 8),
        (2, 13, FOREST, 12, 5),
        (3, 11, CHARCOAL, 9, 4),
    ):
        style = styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(DEEP_GREEN)

    cover_title = styles.add_style("NCIE Cover Title", WD_STYLE_TYPE.PARAGRAPH)
    cover_title.base_style = normal
    cover_title.font.name = "Aptos Display"
    cover_title._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
    cover_title.font.size = Pt(28)
    cover_title.font.bold = True
    cover_title.font.color.rgb = RGBColor.from_string(DEEP_GREEN)
    cover_title.paragraph_format.space_after = Pt(3)
    cover_title.paragraph_format.keep_with_next = True

    front_heading = styles.add_style("NCIE Front Heading", WD_STYLE_TYPE.PARAGRAPH)
    front_heading.base_style = normal
    front_heading.font.name = "Aptos Display"
    front_heading._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
    front_heading.font.size = Pt(18)
    front_heading.font.bold = True
    front_heading.font.color.rgb = RGBColor.from_string(DEEP_GREEN)
    front_heading.paragraph_format.space_after = Pt(8)
    front_heading.paragraph_format.keep_with_next = True

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Aptos"
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    subtitle.font.size = Pt(13)
    subtitle.font.bold = False
    subtitle.font.italic = False
    subtitle.font.color.rgb = RGBColor.from_string(MID_GREY)

    caption = styles["Caption"]
    caption.font.name = "Aptos"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    caption.font.size = Pt(8.5)
    caption.font.bold = False
    caption.font.color.rgb = RGBColor.from_string(MID_GREY)

    bullet = styles["List Bullet"]
    bullet.font.name = "Aptos"
    bullet._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    bullet.font.size = Pt(10.5)
    bullet.font.bold = False
    bullet.font.color.rgb = RGBColor.from_string(CHARCOAL)
    bullet.paragraph_format.space_after = Pt(4)


def configure_section(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)
    section.different_first_page_header_footer = True


def add_header_footer(section):
    header = section.header
    table = header.add_table(rows=1, cols=2, width=Cm(17.6))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_table_borders(table)
    left, right = table.rows[0].cells
    left.width = Cm(12.3)
    right.width = Cm(5.3)
    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("NCIE  |  PRODUCTION DOCUMENTATION SUITE")
    style_run(run, size=8, color=FOREST, bold=True)
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("CONTROLLED REGISTER")
    style_run(run, size=8, color=MID_GREY, bold=False)
    set_paragraph_bottom_border(left.paragraphs[0], color=GOLD, size="8", space="3")
    set_paragraph_bottom_border(right.paragraphs[0], color=GOLD, size="8", space="3")

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    first = p.add_run("NCIE  •  Intelligence  |  Integrity  |  Insight  |  Impact     ")
    style_run(first, size=8, color=MID_GREY, bold=False)
    page_label = p.add_run("Page ")
    style_run(page_label, size=8, color=MID_GREY, bold=False)
    page = add_field(p, "PAGE")
    style_run(page, size=8, color=MID_GREY, bold=False)
    of_label = p.add_run(" of ")
    style_run(of_label, size=8, color=MID_GREY, bold=False)
    pages = add_field(p, "NUMPAGES")
    style_run(pages, size=8, color=MID_GREY, bold=False)


def add_cover(document, total_chapters):
    banner = document.add_table(rows=1, cols=1)
    banner.alignment = WD_TABLE_ALIGNMENT.CENTER
    banner.autofit = False
    remove_table_borders(banner)
    cell = banner.cell(0, 0)
    set_cell_shading(cell, DEEP_GREEN)
    set_cell_margins(cell, top=260, start=180, bottom=260, end=180)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture = p.add_run()
    picture.add_picture(str(LOGO), width=Inches(3.25))
    set_picture_alt_text(
        picture,
        "NCIE logo",
        "National Communications Intelligence Ecosystem emblem in green, gold, white and red.",
    )

    document.add_paragraph()
    title = document.add_paragraph(style="NCIE Cover Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("NCIE → Codex").bold = True
    subtitle = document.add_paragraph(style="NCIE Cover Title")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Production Documentation Suite")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(FOREST)

    rule = document.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_bottom_border(rule, color=GOLD, size="20", space="4")

    descriptor = document.add_paragraph(style="Subtitle")
    descriptor.alignment = WD_ALIGN_PARAGRAPH.CENTER
    descriptor.add_run("Controlled documentation register and production handover index")

    summary = document.add_paragraph(style="NCIE Body")
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = summary.add_run(f"18 documents  •  {total_chapters} chapters")
    style_run(run, size=11, color=CHARCOAL, bold=False)

    document.add_paragraph()
    control = document.add_table(rows=4, cols=2)
    control.alignment = WD_TABLE_ALIGNMENT.CENTER
    control.autofit = False
    remove_table_borders(control)
    values = [
        ("Document", "NCIE Production Documentation Suite"),
        ("Edition", "Production-ready register"),
        ("Prepared", date.today().strftime("%d %B %Y")),
        ("Maintained by", "National Communications Intelligence Ecosystem (NCIE)"),
    ]
    for row, (label, value) in zip(control.rows, values):
        set_cell_width(row.cells[0], Cm(3.7))
        set_cell_width(row.cells[1], Cm(10.5))
        set_cell_shading(row.cells[0], PALE_GREEN)
        set_cell_shading(row.cells[1], "F8FAF8")
        for cell in row.cells:
            set_cell_margins(cell, top=70, start=120, bottom=70, end=120)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        label_p = row.cells[0].paragraphs[0]
        label_run = label_p.add_run(label)
        style_run(label_run, size=8.5, color=FOREST, bold=True)
        value_p = row.cells[1].paragraphs[0]
        value_run = value_p.add_run(value)
        style_run(value_run, size=8.5, color=CHARCOAL, bold=False)

    document.add_page_break()


def add_contents(document):
    document.add_paragraph("Contents", style="NCIE Front Heading")
    intro = document.add_paragraph(style="NCIE Body")
    intro.add_run(
        "This table of contents updates automatically in Microsoft Word. "
        "If page numbers do not appear immediately, select the field and choose Update Field."
    )
    toc = document.add_paragraph()
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u')
    document.add_page_break()


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="NCIE Body")
    paragraph.paragraph_format.left_indent = Cm(0.65)
    paragraph.paragraph_format.first_line_indent = Cm(-0.45)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_together = True
    bullet = paragraph.add_run("•  ")
    style_run(bullet, size=10.5, color=CHARCOAL, bold=False)
    body = paragraph.add_run(text)
    style_run(body, size=10.5, color=CHARCOAL, bold=False)
    return paragraph


def add_overview(document, total_chapters):
    document.add_heading("1. Executive Overview", level=1)
    paragraph = document.add_paragraph(style="NCIE Body")
    paragraph.add_run(
        "The NCIE Production Documentation Suite is the controlled index for eighteen "
        "specifications covering the National Communications Intelligence Ecosystem from "
        "product requirements and architecture through governance, security, delivery, "
        f"verification and production handover. Together, the suite comprises {total_chapters} chapters."
    )

    document.add_heading("1.1 Purpose of this register", level=2)
    for item in (
        "Provide a single, readable reference for the complete documentation suite.",
        "Assign every specification a stable NCIE identifier for traceability.",
        "State the chapter count and primary purpose of each specification.",
        "Support controlled implementation, review, assurance and handover activities.",
    ):
        add_bullet(document, item)

    document.add_heading("1.2 Using the suite", level=2)
    p = document.add_paragraph(style="NCIE Body")
    p.add_run(
        "Use the NCIE document identifier when referencing requirements, decisions, controls "
        "or implementation work. Read the Master Product Requirements Document first, then "
        "use the architecture, governance, engineering, delivery and acceptance specifications "
        "as the authoritative domain references for their stated purposes."
    )


def add_summary(document, documents, total_chapters):
    document.add_heading("2. Suite at a Glance", level=1)
    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    remove_table_borders(table)
    cards = (
        (str(len(documents)), "Controlled documents"),
        (str(total_chapters), "Total chapters"),
        ("NCIE-001—018", "Identifier range"),
    )
    for cell, (metric, label) in zip(table.rows[0].cells, cards):
        set_cell_width(cell, Cm(5.5))
        set_cell_shading(cell, DEEP_GREEN)
        set_cell_margins(cell, top=200, start=120, bottom=200, end=120)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        metric_run = p.add_run(metric)
        style_run(metric_run, size=18, color=GOLD, bold=True)
        p.add_run("\n")
        label_run = p.add_run(label)
        style_run(label_run, size=8.5, color=WHITE, bold=False)

    document.add_paragraph()
    p = document.add_paragraph(style="NCIE Body")
    p.add_run(
        "The register spans product definition, solution and data architecture, artificial "
        "intelligence and memory, governance and security, evidence and integrations, user "
        "experience, functional modules, storage, infrastructure, testing, implementation "
        "planning and final production handover."
    )


def add_register(document, documents):
    document.add_heading("3. Documentation Register", level=1)
    p = document.add_paragraph(style="NCIE Body")
    p.add_run(
        "The following register preserves the identifiers, document titles, chapter counts "
        "and primary purposes defined in the source documentation index."
    )

    table = document.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    widths = (Cm(2.0), Cm(6.5), Cm(1.7), Cm(6.4))
    labels = ("ID", "Document", "Chapters", "Primary purpose")
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for cell, width, label in zip(header.cells, widths, labels):
        set_cell_width(cell, width)
        set_cell_shading(cell, DEEP_GREEN)
        set_cell_margins(cell, top=110, start=100, bottom=110, end=100)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if label in ("ID", "Chapters") else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(label)
        style_run(run, size=8.5, color=WHITE, bold=True)

    for index, item in enumerate(documents, start=1):
        row = table.add_row()
        prevent_row_split(row)
        fill = WHITE if index % 2 else PALE_GREEN
        values = (item[0], item[1], item[2], item[3])
        for column, (cell, width, value) in enumerate(zip(row.cells, widths, values)):
            set_cell_width(cell, width)
            set_cell_shading(cell, fill)
            set_cell_margins(cell, top=90, start=90, bottom=90, end=90)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if column in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(value)
            style_run(
                run,
                size=8.2,
                color=DEEP_GREEN if column == 0 else CHARCOAL,
                bold=True if column == 0 else False,
            )

    caption = document.add_paragraph(style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption.add_run("Table 1. NCIE production documentation register.")


def add_governance(document, source_stated_total, verified_total):
    document.add_heading("4. Production Control Guidance", level=1)
    p = document.add_paragraph(style="NCIE Body")
    p.add_run(
        "Apply the following controls when maintaining or issuing documents in this suite. "
        "These controls support traceability and do not replace formal organizational approval."
    )
    for item in (
        "Retain the assigned NCIE identifier throughout the document lifecycle.",
        "Record version, owner, review status and approval date in each controlled document.",
        "Reference related specifications by identifier and title, not by title alone.",
        "Maintain change history and approval evidence for every released revision.",
        "Validate accessibility, factual integrity, security classification and distribution controls before release.",
    ):
        add_bullet(document, item)

    document.add_heading("4.1 Brand and editorial standard", level=2)
    p = document.add_paragraph(style="NCIE Body")
    p.add_run(
        "Use the NCIE logo and its deep-green, gold, emerald and neutral palette consistently. "
        "Reserve bold type for headings, table headers, identifiers and deliberate emphasis. "
        "Body text beneath headings should remain regular weight for readability."
    )

    document.add_heading("4.2 Source integrity", level=2)
    p = document.add_paragraph(style="NCIE Body")
    p.add_run(
        f"This production-ready register preserves the source catalogue of eighteen documents. "
        f"The source summary stated {source_stated_total} chapters; arithmetic verification of "
        f"the document register produced {verified_total} chapters, which is the total used in "
        "this edition. Operational approvals, signatures, audit evidence, residency controls "
        "and technical enforcement must be recorded by the responsible authority in the relevant "
        "controlled systems; they are not asserted by this document."
    )


def read_source_register():
    source = Document(SOURCE)
    if len(source.tables) != 1:
        raise ValueError("Expected one source register table.")
    rows = source.tables[0].rows
    documents = []
    for row in rows[1:]:
        values = [cell.text.strip() for cell in row.cells]
        if len(values) != 4:
            raise ValueError("Expected four columns in the source register.")
        documents.append(tuple(values))
    total_chapters = sum(int(item[2]) for item in documents)
    source_summary = " ".join(p.text for p in source.paragraphs)
    match = re.search(r"Total:\s*\d+\s*documents\s*/\s*(\d+)\s*chapters", source_summary)
    source_stated_total = int(match.group(1)) if match else total_chapters
    if len(documents) != 18:
        raise ValueError(
            f"Source integrity check failed: documents={len(documents)}, chapters={total_chapters}"
        )
    return documents, total_chapters, source_stated_total


def main():
    documents, total_chapters, source_stated_total = read_source_register()
    document = Document()
    configure_styles(document)
    configure_section(document.sections[0])
    add_header_footer(document.sections[0])
    set_update_fields(document)

    document.core_properties.title = "NCIE → Codex Production Documentation Suite"
    document.core_properties.subject = "Controlled NCIE documentation register"
    document.core_properties.author = "National Communications Intelligence Ecosystem (NCIE)"
    document.core_properties.category = "Production Documentation"
    document.core_properties.keywords = "NCIE, documentation, Codex, architecture, governance, production"
    document.core_properties.comments = (
        "Production-ready layout derived from NCIE Documentations.docx using the NCIE logo palette."
    )

    add_cover(document, total_chapters)
    add_contents(document)
    add_overview(document, total_chapters)
    add_summary(document, documents, total_chapters)
    add_register(document, documents)
    add_governance(document, source_stated_total, total_chapters)

    document.save(OUTPUT)
    print(f"output={OUTPUT}")
    print(f"documents={len(documents)}")
    print(f"chapters={total_chapters}")
    print(f"source_stated_chapters={source_stated_total}")


if __name__ == "__main__":
    main()
