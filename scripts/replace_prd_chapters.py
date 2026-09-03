from copy import deepcopy
from pathlib import Path
import re
import shutil

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

from build_ncie_prd_production import (
    CHARCOAL,
    CHAPTER_MARKER_RE,
    DEEP_GREEN,
    FOREST,
    GOLD,
    LIGHT_GREY,
    NUMBERED_L2_RE,
    NUMBERED_L3_RE,
    PALE_GREEN,
    WHITE,
    collect_diagram_blocks,
    format_diagrams,
    looks_like_short_heading,
    paragraph_all_bold,
    remove_paragraph,
    set_body_all_regular,
    set_heading_runs,
    set_paragraph_left_border,
    set_paragraph_shading,
    validate_output,
)
from build_ncie_production_document import (
    prevent_row_split,
    set_cell_margins,
    set_cell_shading,
    set_repeat_table_header,
    set_table_borders,
)


ROOT = Path(__file__).resolve().parents[1]
EDIT_PATH = ROOT / "Chapters For Edit.docx"
MASTER_PATH = ROOT / "NCIE-001_Master_PRD_v1_2_PRODUCTION_READY.docx"
WORKING_PATH = ROOT / "NCIE-001_Master_PRD_v1_2_CHAPTER_EDIT_WORKING.docx"
BACKUP_PATH = ROOT / "NCIE-001_Master_PRD_v1_2_PRE_CHAPTER_EDIT_BACKUP.docx"

TARGETS = (2, 6, 8, 9, 10, 14, 15, 16, 17, 18)
EDIT_STARTS = {
    2: "Chapter 2",
    6: "Chapter 6",
    8: "Chapter 8",
    9: "Chapter 9",
    10: "Replacement Chapter 10",
    14: "Chapter 14",
    15: "Chapter 15",
    16: "Chapter 16",
    17: "Chapter 17",
    18: "Chapter 18",
}
TITLES = {
    2: "Product Scope, Functional Boundaries & NCIE Intelligence Domain Architecture",
    6: "Data, Intelligence Domain, Canonical Model & Information Governance",
    8: "Intelligence Fusion, Cross-Domain Correlation & National Intelligence Synthesis Governance",
    9: "Regulatory Intelligence, Compliance, Revenue Assurance & Governed Case Progression",
    10: "Situational Awareness, Early Warning, Alerting & National Communications Operating Picture",
    14: "Security, Privacy, Identity, Authorization & Intelligence Sovereignty Governance",
    15: "NCIE Intelligence Assistant, Conversational Intelligence, Collaborative Brainstorming & Governed Memory",
    16: "Data Acquisition, Integration, Source Intelligence & Ingestion Governance",
    17: "Automation Orchestration, RPA Governance, Scheduling & Autonomous Workflow Execution",
    18: "Governance, Auditability, Explainability, Human Oversight, Operational Resilience & Product Acceptance",
}


def element_text(element):
    return "".join(node.text or "" for node in element.iter(qn("w:t"))).strip()


def normalize_space(text):
    return " ".join(text.split())


def set_xml_paragraph_text(element, text):
    p_pr = element.find(qn("w:pPr"))
    for child in list(element):
        if child is not p_pr:
            element.remove(child)
    run = OxmlElement("w:r")
    text_node = OxmlElement("w:t")
    if text[:1].isspace() or text[-1:].isspace():
        text_node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    text_node.text = text
    run.append(text_node)
    element.append(run)


def find_edit_segments(edit_document):
    children = list(edit_document._body._element)
    starts = {}
    for index, child in enumerate(children):
        if child.tag != qn("w:p"):
            continue
        text = normalize_space(element_text(child))
        for number, expected in EDIT_STARTS.items():
            if text == expected and number not in starts:
                starts[number] = index
    missing = set(TARGETS) - set(starts)
    if missing:
        raise ValueError(f"Missing edit chapter starts: {sorted(missing)}")
    ordered = sorted((index, number) for number, index in starts.items())
    segments = {}
    for position, (start, number) in enumerate(ordered):
        end = ordered[position + 1][0] if position + 1 < len(ordered) else len(children)
        while end > start and children[end - 1].tag == qn("w:sectPr"):
            end -= 1
        segments[number] = children[start:end]
    return segments


def master_chapter_ranges(document):
    children = list(document._body._element)
    markers = []
    for index, child in enumerate(children):
        if child.tag != qn("w:p"):
            continue
        text = normalize_space(element_text(child))
        match = CHAPTER_MARKER_RE.fullmatch(text)
        if match:
            markers.append((index, int(match.group(1))))
    first_by_number = {}
    for index, number in markers:
        first_by_number.setdefault(number, index)
    if set(first_by_number) != set(range(1, 19)):
        raise ValueError(f"Master chapter markers invalid: {sorted(first_by_number)}")
    ranges = {}
    ordered = sorted((index, number) for number, index in first_by_number.items())
    for position, (start, number) in enumerate(ordered):
        end = ordered[position + 1][0] if position + 1 < len(ordered) else len(children)
        while end > start and children[end - 1].tag == qn("w:sectPr"):
            end -= 1
        ranges[number] = (start, end)
    return ranges


def merge_numbering(edit_document, master_document, copied_segments):
    used_ids = set()
    for elements in copied_segments.values():
        for element in elements:
            for num_id in element.iter(qn("w:numId")):
                value = num_id.get(qn("w:val"))
                if value is not None:
                    used_ids.add(int(value))

    source_root = edit_document.part.numbering_part.element
    destination_root = master_document.part.numbering_part.element
    source_nums = {
        int(node.get(qn("w:numId"))): node
        for node in source_root.findall(qn("w:num"))
    }
    source_abstracts = {
        int(node.get(qn("w:abstractNumId"))): node
        for node in source_root.findall(qn("w:abstractNum"))
    }
    destination_num_ids = [
        int(node.get(qn("w:numId"))) for node in destination_root.findall(qn("w:num"))
    ]
    destination_abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in destination_root.findall(qn("w:abstractNum"))
    ]
    next_num_id = max(destination_num_ids, default=0) + 1
    next_abstract_id = max(destination_abstract_ids, default=0) + 1
    mapping = {}

    for old_num_id in sorted(used_ids):
        source_num = source_nums.get(old_num_id)
        if source_num is None:
            continue
        source_abstract_id = int(
            source_num.find(qn("w:abstractNumId")).get(qn("w:val"))
        )
        source_abstract = source_abstracts[source_abstract_id]
        abstract_copy = deepcopy(source_abstract)
        abstract_copy.set(qn("w:abstractNumId"), str(next_abstract_id))
        first_num = destination_root.find(qn("w:num"))
        if first_num is None:
            destination_root.append(abstract_copy)
        else:
            destination_root.insert(destination_root.index(first_num), abstract_copy)

        num_copy = deepcopy(source_num)
        num_copy.set(qn("w:numId"), str(next_num_id))
        num_copy.find(qn("w:abstractNumId")).set(qn("w:val"), str(next_abstract_id))
        destination_root.append(num_copy)
        mapping[old_num_id] = next_num_id
        next_num_id += 1
        next_abstract_id += 1

    for elements in copied_segments.values():
        for element in elements:
            for num_id in element.iter(qn("w:numId")):
                value = num_id.get(qn("w:val"))
                if value is not None and int(value) in mapping:
                    num_id.set(qn("w:val"), str(mapping[int(value)]))
    return len(mapping)


def clean_copied_segment(number, source_elements):
    copied = [deepcopy(element) for element in source_elements]
    cleaned = []
    title_continued = normalize_space(TITLES[number] + " — Continued")
    continuation_pattern = re.compile(
        rf"(?:#\s*)?NCIE-001\s*[\u2014-]\s*Replacement\s+Chapter\s+{number}\b.*$",
        re.I,
    )
    first_paragraph = True
    for element in copied:
        if element.tag != qn("w:p"):
            cleaned.append(element)
            continue
        raw_text = element_text(element)
        normalized = normalize_space(raw_text)
        if first_paragraph:
            set_xml_paragraph_text(element, f"Chapter {number} Expansion")
            first_paragraph = False
            cleaned.append(element)
            continue
        if normalized == title_continued:
            continue
        if re.fullmatch(
            rf"NCIE-001\s*[\u2014-]\s*Replacement\s+Chapter\s+{number}",
            normalized,
            re.I,
        ):
            continue
        if continuation_pattern.search(raw_text):
            retained = continuation_pattern.sub("", raw_text)
            retained = re.sub(r"^\s*#{1,4}\s*", "", retained)
            retained = retained.replace("**", "").strip()
            if not retained:
                continue
            set_xml_paragraph_text(element, retained)
        elif "Replacement Chapter 10" in raw_text:
            set_xml_paragraph_text(element, raw_text.replace("Replacement Chapter 10", "Chapter 10"))
        elif re.search(r"(?:^|\s)(?:#{1,4}\s|\*\*)", raw_text):
            cleaned_text = re.sub(r"^\s*#{1,4}\s*", "", raw_text).replace("**", "")
            set_xml_paragraph_text(element, cleaned_text)
        cleaned.append(element)
    return cleaned


def chapter_semantic_text(document, number):
    ranges = master_chapter_ranges(document)
    start, end = ranges[number]
    children = list(document._body._element)[start:end]
    values = []
    for child in children:
        if child.tag == qn("w:p"):
            text = normalize_space(element_text(child))
            if re.match(r"^Figure\s+\d+\.", text, re.I):
                continue
            values.append(text)
        elif child.tag == qn("w:tbl"):
            values.append("[TABLE]" + normalize_space(element_text(child)))
    return "\n".join(values)


def replace_segments(master_document, prepared_segments):
    body = master_document._body._element
    for number in reversed(TARGETS):
        ranges = master_chapter_ranges(master_document)
        start, end = ranges[number]
        children = list(body)
        insertion_reference = children[end] if end < len(children) else None
        for child in children[start:end]:
            body.remove(child)
        if insertion_reference is None:
            insert_at = len(body) - (1 if body.sectPr is not None else 0)
        else:
            insert_at = body.index(insertion_reference)
        for offset, element in enumerate(prepared_segments[number]):
            body.insert(insert_at + offset, element)


def chapter_elements(document, number):
    start, end = master_chapter_ranges(document)[number]
    return list(document._body._element)[start:end]


def format_replacement_paragraphs(document, number):
    elements = chapter_elements(document, number)
    paragraphs = [Paragraph(element, document._body) for element in elements if element.tag == qn("w:p")]
    nonempty = [paragraph for paragraph in paragraphs if paragraph.text.strip()]
    if len(nonempty) < 2:
        raise ValueError(f"Replacement chapter {number} is unexpectedly empty.")
    marker, title = nonempty[0], nonempty[1]
    marker.clear()
    marker.add_run(f"Chapter {number} Expansion")
    marker.style = "NCIE Chapter Label"
    marker.paragraph_format.page_break_before = True
    set_heading_runs(marker)
    title.style = "Heading 1"
    set_heading_runs(title)

    counts = {"body": 0, "heading2": 0, "heading3": 0, "status": 0, "debolded": 0}
    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if not text or paragraph is marker or paragraph is title:
            continue
        original_all_bold = paragraph_all_bold(paragraph)
        if text.startswith("Version:") or text.lower().startswith("chapter status:"):
            paragraph.style = "NCIE Status Block"
            set_paragraph_shading(paragraph, PALE_GREEN)
            set_paragraph_left_border(paragraph)
            set_body_all_regular(paragraph)
            counts["status"] += 1
        elif NUMBERED_L3_RE.match(text):
            paragraph.style = "Heading 3"
            set_heading_runs(paragraph)
            counts["heading3"] += 1
        elif NUMBERED_L2_RE.match(text):
            paragraph.style = "Heading 2"
            set_heading_runs(paragraph)
            counts["heading2"] += 1
        elif looks_like_short_heading(paragraphs, index, paragraph):
            paragraph.style = "Heading 3"
            set_heading_runs(paragraph)
            counts["heading3"] += 1
        else:
            paragraph.style = "NCIE Body"
            if original_all_bold:
                set_body_all_regular(paragraph)
                counts["debolded"] += 1
            else:
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Aptos")
            counts["body"] += 1
    return counts


def format_replacement_tables(document, number):
    tables = [
        Table(element, document._body)
        for element in chapter_elements(document, number)
        if element.tag == qn("w:tbl")
    ]
    for table in tables:
        table.autofit = True
        set_table_borders(table, color=LIGHT_GREY, size="4")
        if not table.rows:
            continue
        set_repeat_table_header(table.rows[0])
        for row_index, row in enumerate(table.rows):
            prevent_row_split(row)
            fill = DEEP_GREEN if row_index == 0 else (WHITE if row_index % 2 else PALE_GREEN)
            for cell in row.cells:
                set_cell_shading(cell, fill)
                set_cell_margins(cell, top=65, start=75, bottom=65, end=75)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(1)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for run in paragraph.runs:
                        run.font.name = "Aptos"
                        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Aptos")
                        run.font.size = Pt(8.2)
                        run.font.color.rgb = RGBColor.from_string(WHITE if row_index == 0 else CHARCOAL)
                        run.bold = row_index == 0
    return len(tables)


def renumber_diagrams(document):
    body_children = list(document._body._element)
    figure_number = 0
    for index, element in enumerate(body_children):
        if element.tag != qn("w:p"):
            continue
        paragraph = Paragraph(element, document._body)
        if paragraph.style.name != "NCIE Figure Caption":
            continue
        next_table = None
        for following in body_children[index + 1:index + 3]:
            if following.tag == qn("w:tbl"):
                next_table = following
                break
            if following.tag == qn("w:p") and element_text(following):
                break
        if next_table is None:
            continue
        tbl_caption = next_table.tblPr.find(qn("w:tblCaption"))
        if tbl_caption is None or not tbl_caption.get(qn("w:val"), "").startswith("NCIE Figure"):
            continue
        figure_number += 1
        old_value = tbl_caption.get(qn("w:val"), "")
        diagram_type = "relationship model" if "relationship" in old_value else "process flow"
        paragraph.clear()
        paragraph.style = "NCIE Figure Caption"
        run = paragraph.add_run(f"Figure {figure_number}. NCIE {diagram_type}.")
        run.bold = False
        run.italic = True
        run.font.color.rgb = RGBColor.from_string(FOREST)
        tbl_caption.set(qn("w:val"), f"NCIE Figure {figure_number}: {diagram_type}")
    return figure_number


def validate_replacements(document, untouched_signatures):
    ranges = master_chapter_ranges(document)
    for number in TARGETS:
        start, end = ranges[number]
        elements = list(document._body._element)[start:end]
        paragraphs = [Paragraph(element, document._body) for element in elements if element.tag == qn("w:p")]
        nonempty = [paragraph for paragraph in paragraphs if paragraph.text.strip()]
        if nonempty[0].text.strip() != f"Chapter {number} Expansion":
            raise ValueError(f"Chapter {number} marker was not normalized.")
        if normalize_space(nonempty[1].text) != TITLES[number]:
            raise ValueError(f"Chapter {number} title mismatch: {nonempty[1].text!r}")
        effective_break = nonempty[0].paragraph_format.page_break_before
        if effective_break is None:
            effective_break = nonempty[0].style.paragraph_format.page_break_before
        if effective_break is not True:
            raise ValueError(f"Chapter {number} does not start on a new page.")
    all_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    if "Replacement Chapter 10" in all_text:
        raise ValueError("The Chapter 10 replacement label remains in the document.")
    if re.search(r"NCIE-001\s*[\u2014-]\s*Replacement Chapter", all_text):
        raise ValueError("Replacement continuation headers remain in the document.")
    for number, signature in untouched_signatures.items():
        if chapter_semantic_text(document, number) != signature:
            raise ValueError(f"Untouched Chapter {number} changed unexpectedly.")


def main():
    edit_document = Document(EDIT_PATH)
    master_document = Document(MASTER_PATH)
    untouched = {
        number: chapter_semantic_text(master_document, number)
        for number in range(1, 19)
        if number not in TARGETS
    }
    source_segments = find_edit_segments(edit_document)
    prepared_segments = {
        number: clean_copied_segment(number, source_segments[number]) for number in TARGETS
    }
    numbering_mappings = merge_numbering(edit_document, master_document, prepared_segments)
    replace_segments(master_document, prepared_segments)

    formatting_counts = {}
    source_table_count = 0
    for number in TARGETS:
        formatting_counts[number] = format_replacement_paragraphs(master_document, number)
        source_table_count += format_replacement_tables(master_document, number)

    allowed_elements = {
        element
        for number in TARGETS
        for element in chapter_elements(master_document, number)
        if element.tag == qn("w:p")
    }
    diagram_blocks = [
        block
        for block in collect_diagram_blocks(master_document)
        if block["paragraphs"]
        and all(paragraph._p in allowed_elements for paragraph in block["paragraphs"])
    ]
    new_diagrams = format_diagrams(master_document, diagram_blocks)
    total_diagrams = renumber_diagrams(master_document)
    validate_replacements(master_document, untouched)

    master_document.core_properties.comments = (
        "Production-ready NCIE-001 v1.2 master updated with authorized complete replacement "
        "chapters 2, 6, 8, 9, 10, and 14–18 from Chapters For Edit.docx. "
        "Chapter 10 replacement labeling was normalized to Chapter 10."
    )
    master_document.save(WORKING_PATH)
    validate_output(WORKING_PATH)

    if not BACKUP_PATH.exists():
        shutil.copy2(MASTER_PATH, BACKUP_PATH)
    shutil.copy2(WORKING_PATH, MASTER_PATH)
    WORKING_PATH.unlink()
    print(f"master={MASTER_PATH}")
    print(f"backup={BACKUP_PATH}")
    print(f"numbering_mappings={numbering_mappings}")
    print(f"replacement_source_tables={source_table_count}")
    print(f"new_diagrams={new_diagrams}")
    print(f"total_diagrams={total_diagrams}")
    print(f"formatting_counts={formatting_counts}")
    print(f"validation={validate_output(MASTER_PATH)}")


if __name__ == "__main__":
    main()
